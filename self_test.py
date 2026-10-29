from core.tz_parser import parse_one, parse_tz, Requirements
from core.selector import build_plan, resolve_platform
from core.defaults import DEFAULT_QX_PRICES, DEFAULT_SETTINGS
from copy import deepcopy
from core.currency import parse_price_input, buffered_rate
from core.project_calc import normalize_arsenal_os_name, arsenal_equipment_default
from core.project_io import dump_project, load_project, validate_quantities
from core.tender_fill import extract_characteristics_docx, build_characteristic_rows, fill_docx_characteristics, requirement_status
from core.defaults import AIO_SPECS
from core.price_import import parse_supplier_price_excel


def main():
    # Regression tests need an unchanging baseline. The user can have a previously
    # downloaded/updated CPU catalog and custom QX prices/settings, which are valid
    # for everyday calculations but must not silently change fixed fixture assertions.
    # This affects only THIS self-test process; the Modern app still loads the live
    # external CPU database and all user's saved prices/settings as before.
    from core import cpu_catalog
    cpu_catalog._external_cache = []
    prices = deepcopy(DEFAULT_QX_PRICES)
    settings = deepcopy(DEFAULT_SETTINGS)

    # ТЗ(3)-style: generic SSD capacity must NOT be silently converted to NVMe/M.2.
    text='''
Наименование характеристики | Значение характеристики | Единица измерения характеристики
Диагональ экрана | ≥ 23 | Дюйм
Объем накопителя SSD | ≥ 480 | Гигабайт
Тип видеокарты | Интегрированная |
Тип оперативной памяти | DDR4 |
Частота процессора базовая | ≥ 3.3 | Гигагерц
Частота процессора максимальная | ≥ 4.3 | Гигагерц
Количество потоков процессора | ≥ 8 | Штука
Количество ядер процессора | ≥ 4 | Штука
Максимальный объем оперативной памяти | ≥ 64 | Гигабайт
Объем установленной оперативной памяти | ≥ 16 | Гигабайт
Тактовая частота оперативной памяти | ≥ 2666 |
Беспроводная связь | Wi-Fi |
'''
    r=parse_one(text)
    assert r.quantity == 0
    assert r.product_type == 'aio24'
    assert r.need_integrated_graphics is True
    assert r.ssd_gb_min == 480 and r.ssd_interface is None and r.ssd_m2_required is False
    rows,reg,warnings,checks=build_plan(r,prices,settings)
    by={x['category']:x for x in rows}
    assert by['SSD']['name'] == 'SSD 480 ГБ', by['SSD']
    assert 'NVMe' not in by['SSD']['name'] and 'M.2' not in by['SSD']['name']
    assert by['Процессор']['name'] == 'AMD Ryzen 5 PRO 5650GE OEM', by['Процессор']['name']
    assert 'AMD Ryzen 5 5500GT' in by['Процессор']['alternatives']
    assert 'AMD Ryzen 5 PRO 5650GE' in by['Процессор']['alternatives']
    assert by['Блок питания']['name'] == 'Встроенный в баребон QX' and by['Блок питания']['price'] == 0

    # v1.9 regression: a table may specify SSD capacity only as total 2.5-inch capacity,
    # while HDD is expressed in TB. These must never fall back to 256 GB SSD / 2 TB HDD.
    drive_tz='''
Наименование характеристики | Значение характеристики | Единица измерения характеристики | Инструкция по заполнению характеристик в заявке
Интерфейс накопителя SSD | SATA | | Значение характеристики не может изменяться участником закупки
Количество накопителей типа HDD | ≥ 1 | Штука | Участник закупки указывает в заявке конкретное значение характеристики
Объем накопителя HDD | ≥ 1 | Терабайт | Участник закупки указывает в заявке конкретное значение характеристики
Общий объем накопителей SSD форм-фактора 2,5" | ≥ 480 | Гигабайт | Участник закупки указывает в заявке конкретное значение характеристики
Количество накопителей типа SSD форм-фактора 2,5" | ≥ 1 | Штука | Участник закупки указывает в заявке конкретное значение характеристики
Общее количество установленных портов DisplayPort | ≥ 1 | Штука | Участник закупки указывает в заявке конкретное значение характеристики
'''
    dr=parse_one(drive_tz)
    assert dr.ssd_gb_min==480 and dr.ssd_25_required is True and dr.ssd_interface=='SATA'
    assert dr.hdd_gb_min==1000 and dr.system_dp_min==1
    drows,dreg,dwarn,dchecks=build_plan(dr,prices,settings)
    dby={x['category']:x for x in drows}
    assert dby['SSD']['name']=='SSD 480 ГБ SATA 2.5"', dby['SSD']
    assert dby['HDD']['name']=='HDD 1 ТБ', dby['HDD']
    assert any(x['Проверка']=='Объём SSD, ГБ' and x['Статус']=='OK' for x in dchecks)
    assert any(x['Проверка']=='Объём HDD, ГБ' and x['Статус']=='OK' for x in dchecks)

    # Current QX AIO platform facts: 23.8 FHD / 27 2K are 75 Hz and always include a
    # non-registry cardreader with 1x Gen1 Type-A + 1x Gen1 Type-C in the platform price.
    assert AIO_SPECS['aio24']['refresh_hz'] == 75 and AIO_SPECS['aio27']['refresh_hz'] == 75
    assert AIO_SPECS['aio24']['built_in_cardreader_ports']['usb3_gen1_type_a'] == 1
    aio_req=Requirements(product_type='aio27', quantity=1, ram_type='DDR5', need_cardreader=True,
                         monitor_refresh_hz_min=100, usb3_gen1_type_c_total_min=1)
    aio_rows,aio_reg,aio_warn,aio_checks=build_plan(aio_req,prices,settings)
    aio_by={x['category']:x for x in aio_rows}
    assert aio_by['Картридер платформы']['price'] == 0 and aio_by['Картридер платформы']['registry'] is False
    assert aio_by['Wi‑Fi']['registry'] is True  # built-in cardreader satisfies the functional cardreader requirement
    assert not any(x.get('name') == 'Картридер (не реестр)' for x in aio_rows)
    pports=aio_by['Порты USB']['constraints']
    assert pports['gen1_c'] >= 1 and pports['gen1_a'] >= 3
    hz_check=next(x for x in aio_checks if x['Проверка']=='Частота обновления экрана, Гц')
    assert hz_check['Факт QX']=='75' and hz_check['Статус']=='НЕ ПРОХОДИТ'

    # Optional front/rear I/O brackets are entered manually and must participate in port totals.
    ext_req=Requirements(product_type='system', quantity=1, platform='amd', usb2_total_min=6, usb2_front_min=4,
                         extra_front_usb2_a=2, extra_rear_gen1_a=2)
    ext_rows,ext_reg,ext_warn,ext_checks=build_plan(ext_req,prices,settings)
    ext_by={x['category']:x for x in ext_rows}
    assert ext_by['Порты USB']['constraints']['usb2_a'] == 6
    assert ext_by['Порты USB']['constraints']['front_usb2'] == 4
    assert ext_by['Порты USB']['constraints']['rear_a'] >= 8
    assert any(x['category']=='Доп. планка портов' and 'передняя панель' in x['name'] for x in ext_rows)
    assert any(x['category']=='Доп. планка портов' and 'задняя панель' in x['name'] for x in ext_rows)
    assert not [x for x in ext_checks if x['Статус']=='НЕ ПРОХОДИТ' and 'USB 2.0' in x['Проверка']]

    # Corporate prose-style TZ: PassMark, I/O, BIOS/manual checks, no wireless, 32 GB dual-channel,
    # SSD of any allowed interface, and quantity 100.
    corp='''
Системный блок Тип 2
Процессор: Passmark CPU на дату подачи заявки не менее 20 000. Поддержка x86 64bit.
Материнская плата: RJ-45 10/100/1000 Мбит/сек. Не менее 1 цифрового порта DisplayPort. Не менее 1 цифрового порта HDMI.
Не менее 1 разъёма PCI-E x1 и x16. Не менее 2-х слотов для оперативной памяти.
Количество USB-портов на передней панели: не менее 2-х USB Type-A, не менее 1-го USB Type-C.
Количество USB-портов на задней панели: не менее 3-х USB Type-A. Макс. поддерживаемый объём оперативной памяти не менее 64GB.
Поддержка одновременной работы не менее 3-х мониторов. Наличие TPM 2.0 или аналогов (Intel PTT, AMD fTPM).
Наличие датчиков температуры, напряжений, оборотов вентилятора, вскрытия.
Обеспечить возможность получения данных о конфигурации BIOS из консоли операционной системы Astra Linux.
Обеспечить возможность изменять пароль администратора BIOS из консоли операционной системы Astra Linux.
Не менее 32GB DDR4/DDR5 (каждая планка памяти не менее 16GB). Частота не менее 2666MHz. Работа в двухканальном режиме.
SSD SATA или М.2 SATA или M.2 PCIe не менее 500GB. Возможность установки дополнительного HDD SATA. TWD не менее 300.
Корпус не более 175 мм х 380 мм х 420 мм. Беспроводная связь: Отсутствует.
Клавиатура проводная полноразмерная с интерфейсом USB Type-A. Мышь проводная оптическая с интерфейсом USB Type-A.
Без ОС. Поддержка Windows 10/11, Astra Linux.
Перечень товаров | Системный блок Тип 2 | 100
'''
    c=parse_one(corp)
    assert c.quantity == 100
    assert c.cpu_passmark_min == 20000 and c.cpu_x86_64_required
    assert c.ram_gb_min == 32 and c.ram_type is None and c.ram_dual_channel_required and c.ram_module_gb_min == 16
    assert c.ssd_gb_min == 500 and c.ssd_interface is None and not c.ssd_m2_required
    assert c.forbid_wireless and not c.need_wifi
    assert c.usb_a_front_min == 2 and c.usb_c_front_min == 1 and c.usb_a_rear_min == 3
    assert c.need_keyboard and c.need_mouse
    rows,reg,warnings,checks=build_plan(c,prices,settings)
    by={x['category']:x for x in rows}
    assert by['Процессор']['name'] == 'AMD Ryzen 7 PRO 5750GE OEM'
    assert by['SSD']['name'] == 'SSD 512 ГБ'
    assert 'Wi‑Fi' not in by
    assert by['Картридер']['registry'] is True
    assert by['Корпус']['name'] == prices['case_protected']['name']
    assert 'Клавиатура + мышь' in by
    assert reg['registry_wifi'] is False and reg['registry_cardreader'] is True

    # Descriptive procurement table: all characteristics are inside one large cell,
    # without KTRU-style "value / instruction" columns. The parser must still see
    # the item as an AIO, quantity, CPU/RAM/SSD/ports/display/power and wireless peripherals.
    from docx import Document as _ProcDoc
    from io import BytesIO as _ProcBIO
    pd=_ProcDoc(); pt=pd.add_table(rows=1, cols=6)
    for i,v in enumerate(['№ п/п','ОКПД2','Наименование','Технические характеристики','Ед. изм.','Кол-во']): pt.rows[0].cells[i].text=v
    pc=pt.add_row().cells
    vals=['1','26.20.15.140','Моноблок',
          'Диагональ экрана (дюйм): не менее 23.8"\n'
          'Разрешение экрана: не менее 1920x1080\n'
          'Технология изготовления матрицы: IPS\n'
          'Контрастность экрана - не менее 1000:1\n'
          'Углы обзора экрана по горизонтали/вертикали - не менее 178/178 градусов\n'
          'Общее количество портов USB2.0 - не менее 2 шт.\n'
          'Общее количество портов USB3.2 Gen1 (USB3.1 Gen1, USB3.0) - не менее 4 шт.\n'
          'Общее количество ядер процессора: не менее 6\n'
          'Максимальное число потоков процессора: не менее 12\n'
          'Базовая тактовая частота - не менее 2.5 ГГц\n'
          'Максимальная тактовая частота - не менее 4.3 ГГц\n'
          'Объем кэш памяти - не менее 12 Мбайт\n'
          'Встроенный графический адаптер – наличие\n'
          'TDP - не более 65 Вт\n'
          'Тип оперативной памяти: не ниже DDR4\n'
          'Объем оперативной памяти: не менее 16 Гигабайт\n'
          'Количество слотов оперативной памяти - не менее 2 шт.\n'
          'Количество разъемов DisplayPort - не менее 1 шт.\n'
          'Количество выходных разъемов HDMI - не менее 1 шт\n'
          'Количество разъемов RJ-45 - не менее 1 шт\n'
          'Общий объем твердотельных накопителей (SSD): Объем SSD накопителя не менее 480 Гигабайт\n'
          'Веб-камера встроенная: наличие\n'
          'Количество мегапикселей веб-камеры: не менее 2 МП\n'
          'Возможность механической блокировки видеопотока камеры - наличие\n'
          'Встроенный микрофон: наличие\n'
          'Встроенные динамики: наличие\n'
          'Встроенный беспроводной модуль Wi-Fi - наличие Встроенный беспроводной модуль Bluetooth – наличие\n'
          'Внешний блок питания – наличие\n'
          'Мощность блока питания - не менее 120 Вт\n'
          'Гарантия: не менее 36 месяцев\n'
          'Общее количество кнопок мыши: не менее 3 (2 + кнопка-колесико)\n'
          'Клавиатура и мышь в комплекте поставки: наличие\n'
          'Тип сенсора мыши: оптический светодиодный\n'
          'Тип подключения мыши: беспроводной\n'
          'Тип подключения клавиатуры: беспроводной\n'
          'Способ нанесения русификации клавиатуры: промышленный;\n'
          'Общее количество клавиш клавиатуры: не менее 104\n'
          'Цифровой блок: наличие\n'
          'Формат клавиатуры: полноразмерная', 'шт.','200']
    for i,v in enumerate(vals): pc[i].text=v
    pd.add_paragraph('Поставляемые моноблоки должны быть совместимы с РЕД ОС, версия 8.')
    pbuf=_ProcBIO(); pd.save(pbuf)
    preqs,_=parse_tz('descriptive.docx',pbuf.getvalue())
    assert len(preqs)==1
    pr=preqs[0]
    assert pr.title=='Моноблок' and pr.product_type=='aio24' and pr.quantity==200
    assert (pr.cpu_cores_min,pr.cpu_threads_min,pr.cpu_freq_ghz_min,pr.cpu_freq_max_ghz_min,pr.cpu_l3_mb_min,pr.cpu_tdp_w_max)==(6,12,2.5,4.3,12,65)
    assert pr.need_integrated_graphics and pr.ram_type is None and pr.ram_type_min_generation==4 and pr.ram_gb_min==16 and pr.ram_slots_min==2
    assert pr.ssd_gb_min==480 and pr.usb2_total_min==2 and pr.usb3_gen1_total_min==4
    assert pr.system_dp_min==1 and pr.system_hdmi_min==1 and pr.ethernet_ports_min==1
    assert pr.monitor_inches==23.8 and pr.monitor_resolution=='1920x1080' and pr.monitor_matrix=='IPS'
    assert pr.monitor_contrast_min==1000 and pr.monitor_view_angle_h_min==178 and pr.monitor_view_angle_v_min==178
    assert pr.webcam_mp_min==2 and pr.webcam_privacy_shutter and pr.need_builtin_mic and pr.monitor_speakers
    assert pr.need_wifi and pr.need_bluetooth and pr.external_psu_required and pr.psu_w_min==120 and pr.warranty_months_min==36
    assert pr.need_keyboard and pr.need_mouse and pr.keyboard_wireless and pr.mouse_wireless
    assert pr.keyboard_keys_min==104 and pr.keyboard_numeric_keypad and pr.keyboard_fullsize and pr.mouse_optical and pr.mouse_buttons_min==3
    prows,preg,pwarn,pchecks=build_plan(pr,prices,settings)
    pby={x['category']:x for x in prows}
    assert pby['Процессор']['name']=='AMD Ryzen 5 PRO 5650GE OEM'
    assert pby['SSD']['name']=='SSD 480 ГБ'
    assert pby['Клавиатура + мышь']['source']=='QX', pby['Клавиатура + мышь']
    # Even a special peripheral TZ must not silently replace QX. The user can opt out.
    pr.keyboard_source_mode='manual'; pr.mouse_source_mode='manual'
    mrows,_,_,_=build_plan(pr,prices,settings)
    assert any(x['category']=='Клавиатура' and x['source']=='Вручную' for x in mrows)
    assert any(x['category']=='Мышь' and x['source']=='Вручную' for x in mrows)
    assert any('внешний блок питания' in w.lower() for w in pwarn)
    assert any(c['Проверка']=='Тип блока питания моноблока' and c['Статус']=='НЕ ПРОХОДИТ' for c in pchecks)


    # Characteristic-table quantity must not be mistaken for procurement quantity.
    qfalse=parse_one('''
Наименование характеристики | Значение характеристики | Единица измерения характеристики | Инструкция по заполнению характеристик в заявке
Количество встроенных в корпус портов USB 3.2 Gen 2 Type-C | ≥ 1 | Штука | Участник закупки указывает в заявке конкретное значение характеристики
''')
    assert qfalse.quantity == 0

    # DOCX application filling: mutable rows get facts, locked rows stay byte-level logically unchanged.
    from docx import Document
    from io import BytesIO
    d=Document(); t=d.add_table(rows=1, cols=4)
    for i,v in enumerate(['Наименование характеристики','Значение характеристики','Единица измерения характеристики','Инструкция по заполнению характеристик в заявке']): t.rows[0].cells[i].text=v
    for vals in [
        ('Количество ядер процессора','≥ 10','Штука','Участник закупки указывает в заявке конкретное значение характеристики'),
        ('Тип оперативной памяти','DDR5','','Значение характеристики не может изменяться участником закупки'),
        ('Количество встроенных в корпус портов USB 3.2 Gen 2 Type-C','≥ 1','Штука','Участник закупки указывает в заявке конкретное значение характеристики'),
    ]:
        cells=t.add_row().cells
        for i,v in enumerate(vals): cells[i].text=v
    bio=BytesIO(); d.save(bio); src_docx=bio.getvalue()
    chars=extract_characteristics_docx(src_docx)
    assert len(chars)==3 and sum(1 for x in chars if x['editable'])==2
    # Reuse the corporate plan above (CPU and USB inventory are known).
    app_rows=build_characteristic_rows(chars,[c.to_dict()],[{'req':c.to_dict(),'rows':rows,'registry':reg,'warnings':warnings,'checks':checks}],{})
    core_row=next(x for x in app_rows if x['name']=='Количество ядер процессора')
    assert core_row['auto_value']=='8' and core_row['value']=='' and core_row['status']=='АВТО НЕ СООТВЕТСТВУЕТ'
    locked=next(x for x in app_rows if x['name']=='Тип оперативной памяти')
    assert locked['value']=='DDR5' and locked['status']=='НЕ ИЗМЕНЯЕТСЯ'
    # Manually provide a compliant value for the rejected CPU fact before Word fill.
    core_override={core_row['id']:'10'}
    app_rows_manual=build_characteristic_rows(chars,[c.to_dict()],[{'req':c.to_dict(),'rows':rows,'registry':reg,'warnings':warnings,'checks':checks}],core_override)
    vals={x['id']:x['value'] for x in app_rows_manual if x['editable'] and x['value']}
    filled=fill_docx_characteristics(src_docx,chars,vals)
    fd=Document(BytesIO(filled)); ft=fd.tables[0]
    assert ft.rows[2].cells[1].text=='DDR5'  # locked row was untouched
    assert ft.rows[1].cells[1].text != '≥ 10'

    # AIO factual export knows the confirmed 75 Hz value.
    d2=Document(); t2=d2.add_table(rows=1, cols=4)
    for i,v in enumerate(['Наименование характеристики','Значение характеристики','Единица измерения характеристики','Инструкция по заполнению характеристик в заявке']): t2.rows[0].cells[i].text=v
    cells=t2.add_row().cells
    for i,v in enumerate(('Частота обновления экрана','≥ 60','Герц','Участник закупки указывает в заявке конкретное значение характеристики')): cells[i].text=v
    bio2=BytesIO(); d2.save(bio2); chars2=extract_characteristics_docx(bio2.getvalue())
    aio_app=build_characteristic_rows(chars2,[aio_req.to_dict()],[{'req':aio_req.to_dict(),'rows':aio_rows,'registry':aio_reg,'warnings':aio_warn,'checks':aio_checks}],{})
    assert aio_app[0]['value']=='75' and aio_app[0]['status']=='OK'

    # A known automatic value that violates the customer's sign must NOT be silently
    # inserted. It is shown as an auto candidate; manual entry is still allowed and warned.
    d3=Document(); t3=d3.add_table(rows=1, cols=4)
    for i,v in enumerate(['Наименование характеристики','Значение характеристики','Единица измерения характеристики','Инструкция по заполнению характеристик в заявке']): t3.rows[0].cells[i].text=v
    cells=t3.add_row().cells
    for i,v in enumerate(('Частота обновления экрана','≥ 100','Герц','Участник закупки указывает в заявке конкретное значение характеристики')): cells[i].text=v
    bio3=BytesIO(); d3.save(bio3); chars3=extract_characteristics_docx(bio3.getvalue())
    rejected=build_characteristic_rows(chars3,[aio_req.to_dict()],[{'req':aio_req.to_dict(),'rows':aio_rows,'registry':aio_reg,'warnings':aio_warn,'checks':aio_checks}],{})
    assert rejected[0]['auto_value']=='75' and rejected[0]['value']=='' and rejected[0]['status']=='АВТО НЕ СООТВЕТСТВУЕТ'
    manual_bad=build_characteristic_rows(chars3,[aio_req.to_dict()],[{'req':aio_req.to_dict(),'rows':aio_rows,'registry':aio_reg,'warnings':aio_warn,'checks':aio_checks}],{chars3[0]['id']:'75'})
    assert manual_bad[0]['value']=='75' and manual_bad[0]['status']=='НЕ СООТВЕТСТВУЕТ'


    # Strict comparison signs from the customer's TZ must be respected exactly.
    assert requirement_status('> 100', '100') == 'НЕ СООТВЕТСТВУЕТ'
    assert requirement_status('> 100', '100.01') == 'OK'
    assert requirement_status('≥ 100', '100') == 'OK'
    assert requirement_status('< 100', '100') == 'НЕ СООТВЕТСТВУЕТ'
    assert requirement_status('≤ 100', '100') == 'OK'
    assert requirement_status('≥ 250 и < 300', '299.9') == 'OK'
    assert requirement_status('≥ 250 и < 300', '300') == 'НЕ СООТВЕТСТВУЕТ'
    assert requirement_status('не менее 20 000', '20000') == 'OK'
    assert requirement_status('не более 90', '91') == 'НЕ СООТВЕТСТВУЕТ'
    assert requirement_status('> 1920 x 1080', '1920 x 1080') == 'НЕ СООТВЕТСТВУЕТ'
    assert requirement_status('≥ 1920 x 1080', '2560 x 1440') == 'OK'

    # A different AIO platform must not inherit QX 23.8/27 facts. All platform facts
    # come only from the specialist's manual inputs and still participate in warnings.
    custom_req=Requirements(product_type='aio27', quantity=1, ram_type='DDR5', aio_platform_mode='custom',
                            custom_aio_name='Custom 27 test', custom_aio_diagonal_in=27,
                            custom_aio_resolution='2560x1440', custom_aio_refresh_hz=100,
                            custom_aio_built_in_cardreader=True, custom_aio_integrated_psu=True,
                            custom_aio_psu_w=180, custom_aio_gen1_a=2, custom_aio_gen1_c=1,
                            custom_aio_front_a=2, custom_aio_front_c=1,
                            monitor_refresh_hz_min=100, usb3_gen1_type_c_total_min=1)
    custom_rows,custom_reg,custom_warn,custom_checks=build_plan(custom_req,prices,settings)
    custom_by={x['category']:x for x in custom_rows}
    assert custom_by['Моноблочная платформа']['name']=='Custom 27 test'
    assert custom_by['Моноблочная платформа']['source']=='Вручную'
    assert custom_by['Блок питания']['constraints']['w']==180
    assert custom_by['Порты USB']['constraints']['gen1_c'] >= 1
    assert any('другая моноблочная платформа' in w.lower() for w in custom_warn)

    # If a custom platform has no factual display values, nothing from QX may leak in.
    unknown_custom=Requirements(product_type='aio27', quantity=1, ram_type='DDR5', aio_platform_mode='custom')
    uc_rows,uc_reg,uc_warn,uc_checks=build_plan(unknown_custom,prices,settings)
    from docx import Document as _Doc
    from io import BytesIO as _BIO
    _d=_Doc(); _t=_d.add_table(rows=1,cols=4)
    for i,v in enumerate(['Наименование характеристики','Значение характеристики','Единица измерения характеристики','Инструкция по заполнению характеристик в заявке']): _t.rows[0].cells[i].text=v
    _r=_t.add_row().cells
    for i,v in enumerate(('Частота обновления экрана','≥ 75','Герц','Участник закупки указывает в заявке конкретное значение характеристики')): _r[i].text=v
    _b=_BIO(); _d.save(_b); _chars=extract_characteristics_docx(_b.getvalue())
    _app=build_characteristic_rows(_chars,[unknown_custom.to_dict()],[{'req':unknown_custom.to_dict(),'rows':uc_rows,'registry':uc_reg,'warnings':uc_warn,'checks':uc_checks}],{})
    assert _app[0]['value']=='' and _app[0]['status']=='НУЖНО ЗАПОЛНИТЬ'

    # Quantity gate.
    assert validate_quantities([r.to_dict()])
    rr=r.to_dict(); rr['quantity']=1
    assert validate_quantities([rr]) == []

    # USD manual-price conversion uses CBR rate + 5 RUB reserve.
    rub,currency,note=parse_price_input('$100',86.8872,5)
    assert currency == 'USD' and abs(rub-9188.72) < 0.001 and '+ 5' in note
    rub2,currency2,_=parse_price_input('15 500',86.8872,5)
    assert rub2 == 15500 and currency2 == 'RUB'
    assert abs(buffered_rate(86.8872,5)-91.8872) < 1e-9

    # Backup round-trip restores the full editable calculation snapshot.
    reqs=[rr]
    plans=[{'req':rr,'rows':rows,'registry':reg,'warnings':warnings,'checks':checks}]
    raw=dump_project(reqs=reqs,plans=plans,source_text='TZ',qx_prices=prices,project_meta={'delivery':30000,'sender':'ЗИОН'},source_filename='TZ.docx',source_docx_bytes=src_docx,tender_overrides={'t0r1':'10'},tender_characteristics=chars)
    restored=load_project(raw)
    assert restored['reqs'][0]['quantity'] == 1
    assert restored['plans'][0]['rows'][0]['name'] == rows[0]['name']
    assert restored['project_meta']['delivery'] == 30000
    assert restored['source_filename']=='TZ.docx' and restored['source_docx_bytes'] == src_docx
    assert restored['tender_overrides']['t0r1']=='10' and len(restored['tender_characteristics'])==3


    # Excel export must tolerate structured requirement values (dict/list) and never crash.
    from core.exporter import export_xlsx
    from openpyxl import load_workbook
    from io import BytesIO
    excel_req = dict(rr)
    excel_req['structured_test'] = {'ports': {'usb_c': 2}, 'options': ['A', 'B']}
    excel_req['manual_requirements'] = ['Проверить BIOS', 'Проверить датчики']
    blob = export_xlsx('export-test', excel_req, rows, reg, {'delivery': 0, 'bg1': 0, 'bg2': 0, 'markup': 0.2}, checks)
    wb = load_workbook(BytesIO(blob), read_only=True)
    assert 'Распознано из ТЗ' in wb.sheetnames and 'Сборка' in wb.sheetnames
    values = [r[1].value for r in wb['Распознано из ТЗ'].iter_rows(min_row=2) if r[0].value == 'structured_test']
    assert values and 'usb_c' in values[0]


    # TZ(5)-style CPU: i5-12600K is the minimal confirmed LGA1700 match.
    cpu_tz = Requirements(product_type='system', quantity=1, platform='intel', ram_type='DDR5',
                          cpu_cores_min=10, cpu_threads_min=16, cpu_freq_ghz_min=3.7,
                          cpu_freq_max_ghz_min=4.2, cpu_l3_mb_min=20, cpu_tdp_w_max=150,
                          need_integrated_graphics=True)
    cpu_rows,cpu_reg,cpu_warn,cpu_checks=build_plan(cpu_tz,prices,settings)
    cpu_row=next(x for x in cpu_rows if x['category']=='Процессор')
    assert cpu_row['name']=='Intel Core i5-12600K OEM', cpu_row
    assert cpu_row['alternatives']==['Intel Core i5-12600K']
    assert not [x for x in cpu_checks if x.get('Статус')=='НЕ ПРОХОДИТ' and ('CPU' in x.get('Проверка','') or 'процессора' in x.get('Проверка',''))]

    # Commercial proposal: outgoing number, markup, model/registry/quantity and totals.
    from core.commercial_offer import build_commercial_offer, calculate_offer_price, outgoing_number
    from datetime import date
    from docx import Document as _KpDoc
    from io import BytesIO as _KpBio
    assert outgoing_number(date(2026,9,15),1)=='БрП-150926-1'
    assert abs(calculate_offer_price(100000,20,'markup')-120000)<0.001
    assert abs(calculate_offer_price(100000,20,'margin')-125000)<0.001
    kp_blob=build_commercial_offer([
        {'name':'Моноблок QX "AIO А520/24-0102" 47237353.466369.002-16',
         'registry_number':'10598727','quantity':2,'unit_price':72486,'total':144972}
    ],sequence=1,offer_date=date(2026,9,15),warranty_months=12,validity_days=30)
    kp_doc=_KpDoc(_KpBio(kp_blob))
    assert any('БрП-150926-1' in p.text for p in kp_doc.paragraphs)
    assert any('15.09.2026' in p.text for p in kp_doc.paragraphs)
    kp_table=kp_doc.tables[1]
    assert '10598727' in kp_table.rows[1].cells[1].text
    assert '72 486' in kp_table.rows[1].cells[3].text
    assert '144 972' in kp_table.rows[-1].cells[-1].text


    # v1.7: a system block may use a completely different chassis. Its price/facts
    # must survive into the plan and participate in USB/TZ checks without inheriting QX case facts.
    _custom_case_req = Requirements(
        product_type='system', quantity=1, platform='amd', cpu_cores_min=6, cpu_threads_min=12,
        system_case_mode='custom', custom_case_name='CUSTOM CASE QA', custom_case_price=3210,
        custom_case_width_mm=190, custom_case_length_mm=410, custom_case_height_mm=430,
        custom_case_bay_525=2, custom_case_bay_25=3, custom_case_bay_35=2,
        custom_case_usb2_front=2, custom_case_gen1_a_front=1, custom_case_gen1_c_front=1,
        usb_c_front_min=1,
    )
    _cc_rows, _cc_rec, _cc_warn, _cc_checks = build_plan(_custom_case_req, prices, settings)
    _cc = next(x for x in _cc_rows if x['category'] == 'Корпус')
    assert _cc['name'] == 'CUSTOM CASE QA' and float(_cc['price']) == 3210
    assert _cc['constraints']['custom'] is True and _cc['constraints']['width_mm'] == 190
    assert any(x['Проверка'] == 'USB Type-C спереди' and x['Статус'] == 'OK' for x in _cc_checks)

    # v2.4: a saved non-QX equipment profile is authoritative. Its motherboard,
    # registry components, prices and board facts must be used without borrowing QX.
    _vendor_profile = {
        'profile_name':'Бештау QA','manufacturer':'Бештау',
        'boards':{'amd':{
            'name':'Бештау AM4 QA','price':7777,'registry':True,'registry_number':'REG-MB-1',
            'socket':'AM4','chipset':'B550','form_factor':'mATX','ram_type':'DDR4',
            'ram_max_gb':64,'ram_max_mhz':3200,'ram_slots':4,'ethernet_mbps':1000,'ethernet_ports':1,
            'm2_key_m':1,'m2_key_e':1,'sata':4,'pcie_x1':1,'pcie_x16':1,'tpm_header':1,
            'video':{'HDMI':1,'DisplayPort':1},
            'usb_external':{'USB 2.0 Type-A':4,'USB 3.2 Gen1 Type-A':2}
        }},
        'components':[
            {'role':'hub','condition':'always','category':'USB-хаб','name':'Бештау HUB QA','qty':1,'price':1234,'registry':True,'registry_number':'REG-HUB-1','supplier':'ExeGate','supplier_sku':'EX-HUB-QA','usb3_gen1_type_a':2,'location':'front'},
            {'role':'cardreader','condition':'cardreader','category':'Картридер','name':'Бештау CR QA','qty':1,'price':800,'registry':True,'registry_number':'REG-CR-1'},
        ],
    }
    _vendor_req = Requirements(product_type='system', quantity=1, platform='amd', equipment_profile_name='Бештау QA',
                               cpu_cores_min=6, ram_gb_min=16, ram_max_supported_gb_min=32, need_cardreader=True)
    _vr, _vreg, _vw, _vc = build_plan(_vendor_req, prices, settings, equipment_profile=_vendor_profile)
    _vby = {x['category']:x for x in _vr}
    assert _vby['Материнская плата']['name']=='Бештау AM4 QA' and _vby['Материнская плата']['price']==7777
    assert _vby['Материнская плата']['source']=='Бештау' and _vby['Материнская плата']['registry'] is True
    assert _vby['USB-хаб']['name']=='Бештау HUB QA' and _vby['USB-хаб']['price']==1234
    assert _vby['USB-хаб']['constraints']['supplier']=='ExeGate' and _vby['USB-хаб']['constraints']['supplier_sku']=='EX-HUB-QA'
    assert _vby['Картридер']['name']=='Бештау CR QA' and _vby['Картридер']['price']==800
    assert not any(x['category']=='USB-хаб' and x['source']=='QX' for x in _vr)
    _chars=[{'id':'vendor-rammax','plan_index':0,'name':'Максимальный объем оперативной памяти','original_value':'≥ 32','unit':'Гигабайт','instruction':'Участник закупки указывает в заявке конкретное значение характеристики','editable':True}]
    _crow=build_characteristic_rows(_chars,[_vendor_req.to_dict()],[{'req':_vendor_req.to_dict(),'rows':_vr,'registry':_vreg,'warnings':_vw,'checks':_vc}],{})[0]
    assert str(_crow['auto_value'])=='64' and 'выбранной базы' in _crow['source']

    # v2.4.1: Arsenal must never display a CPU as the equipment fallback, and
    # a boolean OS presence value must never become the OS product name.
    _ars = arsenal_equipment_default({'req': _vendor_req.to_dict(), 'rows': _vr, 'registry': None})
    assert _ars['equipment'].startswith('Бештау QA') and 'Ryzen' not in _ars['equipment'] and _ars['registry_number']==''
    assert normalize_arsenal_os_name(True, 'Да') == ''
    assert normalize_arsenal_os_name(True, 'Windows 11 Pro') == 'Windows 11 Pro'
    assert normalize_arsenal_os_name(False, None) == 'нет'
    _os_presence = parse_one('Наличие предустановленной операционной системы | Да')
    assert _os_presence.need_os is True and _os_presence.os_text is None

    # v2.5: supplier-price parser auto-detects an ExeGate-style header and reads USD OPT prices.
    from openpyxl import Workbook
    from io import BytesIO
    _wb=Workbook(); _ws=_wb.active; _ws.title='Прайс-лист на 16.09.2026'
    _headers=['Свободный\nостаток, шт.','Резерв,\nшт.','Ожидаемая дата прихода','Код','P/N','Наименование','Фото','Ссылка','Шт. в упаковке*','Цена упаковки*','ОПТ','РРЦ']
    for _c,_v in enumerate(_headers,1): _ws.cell(5,_c).value=_v
    _ws.cell(6,1).value=10; _ws.cell(6,4).value='276689'; _ws.cell(6,5).value='EX276689RUS'; _ws.cell(6,6).value='SSD 480GB'; _ws.cell(6,11).value=63.99; _ws.cell(6,12).value=69.75
    _bio=BytesIO(); _wb.save(_bio)
    _pdata=parse_supplier_price_excel(_bio.getvalue(),'ExeGate','ОПТ')
    assert _pdata['item_count']==1 and _pdata['price_date']=='16.09.2026'
    assert _pdata['items'][0]['part_number']=='EX276689RUS' and _pdata['items'][0]['selected_price_usd']==63.99


    # v2.6 regression: "each installed RAM module >= 8 GB" does NOT mean that
    # two modules are required. With total installed RAM >= 8 GB and no explicit
    # dual-channel / module-count requirement, the minimal valid configuration is 1x8 GB.
    _ram8_text = """
Наименование характеристики | Значение характеристики | Единица измерения характеристики
Объем установленной оперативной памяти | ≥ 8 | Гигабайт
Объем каждого установленного модуля оперативной памяти | ≥ 8 | Гигабайт
Допустимый максимальный объем увеличения оперативной памяти | ≥ 64 | Гигабайт
Тип оперативной памяти | DDR4 |
Тактовая частота оперативной памяти | ≥ 3200 | Мегагерц
"""
    _ram8_req = parse_one(_ram8_text)
    assert _ram8_req.ram_gb_min == 8 and _ram8_req.ram_module_gb_min == 8
    assert _ram8_req.ram_dual_channel_required is False and _ram8_req.ram_max_supported_gb_min == 64
    _ram8_rows, _ram8_reg, _ram8_warn, _ram8_checks = build_plan(_ram8_req, prices, settings)
    _ram8 = next(x for x in _ram8_rows if x['category'] == 'Оперативная память')
    assert _ram8['constraints']['modules'] == 1 and _ram8['constraints']['module_gb'] == 8, _ram8
    assert _ram8['constraints']['min_gb'] == 8 and '16 ГБ' not in _ram8['name'], _ram8

    # The old 2-module behavior remains only when the TZ explicitly requires dual channel.
    _ram_dual = Requirements(product_type='system', quantity=1, ram_type='DDR4', ram_gb_min=8,
                             ram_module_gb_min=8, ram_dual_channel_required=True)
    _dual_rows,_,_,_=build_plan(_ram_dual, prices, settings)
    _dual = next(x for x in _dual_rows if x['category']=='Оперативная память')
    assert _dual['constraints']['modules'] == 2 and _dual['constraints']['module_gb'] == 8
    assert _dual['constraints']['min_gb'] == 16

    # v2.6 peripheral catalogs: a saved monitor/keyboard/mouse is selected only if all
    # parsed requirements it can prove are satisfied. Registry number and price follow the model.
    _periph_catalogs = {
        'monitor':[{'name':'MON QA','price':12345,'registry_number':'MON-REG-1','inches':23.8,
                    'resolution':'1920x1080','refresh_hz':75,'matrix':'IPS','brightness':300,
                    'contrast':1000,'dynamic_contrast':5000000,'response_ms':5,'view_angle_h':178,'view_angle_v':178,
                    'color_depth_bits':8,'aspect_ratio':'16:9','energy_class':'A+','interfaces':['VGA','DisplayPort','HDMI'],
                    'hdmi_ports':1,'dp_ports':1,'speakers':True,'external_psu':True,'wall_mount':True,
                    'vesa':'75 x 75','pivot':False,'tilt':True,'webcam':False,'touch':False,'usb_hub':False,
                    'curved':False,'swivel':False,'cable_included':True,'cable_type':'HDMI-HDMI'}],
        'keyboard':[{'name':'KB QA','price':900,'registry_number':'KB-REG-1','wired':True,'wireless':False,'usb':True,
                    'layout':'ЙЦУКЕН','russification':'Промышленный','flexible':False,'has_usb_port':False,
                    'multifunction':False,'touchpad':False,'palmrest':False,'scroll_wheel':False,
                    'water_protection':False,'detachable_cable':False,'volume_control':False,'backlight':False}],
        'mouse':[{'name':'MOUSE QA','price':700,'registry_number':'MOUSE-REG-1','wired':True,'wireless':False,'usb':True,
                 'optical':True,'dpi':1000,'cable_m':1.5,'programmable_buttons':False,
                 'side_buttons':False,'silent_click':False,'detachable_cable':False}],
    }
    _mr = Requirements(product_type='monitor', quantity=1, monitor_inches=23, monitor_resolution='1920x1080',
                       monitor_refresh_hz_min=75, monitor_matrix='IPS', monitor_brightness_min=300,
                       monitor_brightness_max=350, monitor_contrast_min=1000, monitor_dynamic_contrast_max=10000000,
                       monitor_response_ms_max=6, monitor_view_angle_h_min=178, monitor_view_angle_v_min=178,
                       monitor_color_depth_bits=8, monitor_aspect_ratio='16:9', monitor_energy_class='A+',
                       monitor_interfaces='VGA DisplayPort HDMI', monitor_hdmi_ports_min=1, monitor_dp_ports_min=1,
                       monitor_speakers=True, monitor_external_psu=True, monitor_wall_mount=True, monitor_vesa='VESA 75 x 75',
                       monitor_pivot=False, monitor_tilt=True, monitor_webcam=False, monitor_touch=False,
                       monitor_usb_hub=False, monitor_curved=False, monitor_swivel=False, monitor_cable_included=True,
                       monitor_cable_type='HDMI-HDMI')
    _mrows,_mreg,_mw,_mc = build_plan(_mr, prices, settings, peripheral_catalogs=_periph_catalogs)
    assert _mrows[0]['name']=='MON QA' and float(_mrows[0]['price'])==12345 and _mreg['registry_number']=='MON-REG-1'
    assert not [x for x in _mc if x.get('Статус')=='НЕ ПРОХОДИТ']

    _kr = Requirements(product_type='keyboard', quantity=1, keyboard_source_mode='catalog', keyboard_wired=True, keyboard_usb=True,
                       keyboard_layout='ЙЦУКЕН', keyboard_russification='Промышленный', keyboard_flexible=False,
                       keyboard_has_usb_port=False, keyboard_multifunction=False, keyboard_touchpad=False,
                       keyboard_palmrest=False, keyboard_scroll_wheel=False, keyboard_water_protection=False,
                       keyboard_detachable_cable=False, keyboard_volume_control=False, keyboard_backlight=False)
    _krows,_kreg,_,_ = build_plan(_kr, prices, settings, peripheral_catalogs=_periph_catalogs)
    assert _krows[0]['name']=='KB QA' and _kreg['registry_number']=='KB-REG-1'

    _xr = Requirements(product_type='mouse', quantity=1, mouse_source_mode='catalog', mouse_wired=True, mouse_usb=True, mouse_optical=True,
                       mouse_cable_min_m=1.5, mouse_cable_max_m=2, mouse_dpi_min=1000,
                       mouse_programmable_buttons=False, mouse_side_buttons=False,
                       mouse_silent_click=False, mouse_detachable_cable=False)
    _xrows,_xreg,_,_ = build_plan(_xr, prices, settings, peripheral_catalogs=_periph_catalogs)
    assert _xrows[0]['name']=='MOUSE QA' and _xreg['registry_number']=='MOUSE-REG-1'

    # v2.6.3 regression: merely having saved keyboard/mouse records must NOT replace QX.
    _qx_per = Requirements(product_type='system', quantity=1, need_keyboard=True, need_mouse=True)
    _qx_rows,_,_,_=build_plan(_qx_per,prices,settings,peripheral_catalogs=_periph_catalogs)
    _qx_pair = next(x for x in _qx_rows if x['category']=='Клавиатура + мышь')
    assert _qx_pair['source']=='QX' and 'QX' in _qx_pair['name'], _qx_pair
    _qx_per.keyboard_source_mode='catalog'; _qx_per.mouse_source_mode='catalog'
    _cat_rows,_,_,_=build_plan(_qx_per,prices,settings,peripheral_catalogs=_periph_catalogs)
    assert any(x['name']=='KB QA' for x in _cat_rows) and any(x['name']=='MOUSE QA' for x in _cat_rows)

    # v2.6.1 regression: exact DDR4 must not force AMD when the CPU constraints have
    # no confirmed AM4 APU. The resolver must consider both CPU availability and board
    # compatibility, so a tender like TZ(9) exposes Intel CPU options instead of an
    # unexplained empty selector. The QX Intel DDR5 board mismatch remains visible.
    _cpu_switch = Requirements(product_type='system', quantity=1, ram_type='DDR4', ram_gb_min=16,
                               cpu_cores_min=4, cpu_threads_min=8, cpu_freq_ghz_min=2.5,
                               cpu_freq_max_ghz_min=4.2, cpu_l3_mb_min=20, need_integrated_graphics=True)
    _plat,_pwarn = resolve_platform(_cpu_switch)
    assert _plat == 'intel', (_plat, _pwarn)
    _srows,_,_swarn,_schecks = build_plan(_cpu_switch, prices, settings)
    _scpu = next(x for x in _srows if x['category']=='Процессор')
    assert _scpu['name'] == 'Intel Core i5-13400 OEM' and _scpu['alternatives'], _scpu
    assert any(x['Проверка']=='Тип ОЗУ' and x['Статус']=='НЕ ПРОХОДИТ' for x in _schecks)

    # v2.8 regression: an explicit AMD/Intel selection is authoritative, and standard
    # QX AMD always includes the confirmed Key E adapter. The A520 board has no native
    # M.2 Key E, but the adapter contributes one Key E connection for Wi-Fi and therefore
    # satisfies Key E >=1 without forcing an Intel switch.
    _manual_amd = Requirements(product_type='system', quantity=20, platform='amd', ram_type='DDR4', ram_gb_min=8,
                               ram_mhz_min=3200, cpu_cores_min=4, cpu_threads_min=8, cpu_freq_ghz_min=3.6,
                               cpu_l3_mb_min=4, cpu_tdp_w_max=90, need_integrated_graphics=True,
                               m2_key_m_min=1, m2_key_e_min=1, psu_w_min=400)
    _manual_plat,_manual_warn = resolve_platform(_manual_amd)
    assert _manual_plat == 'amd', (_manual_plat, _manual_warn)
    _manual_rows,_,_manual_plan_warn,_manual_checks = build_plan(_manual_amd, prices, settings)
    _manual_cpu = next(x for x in _manual_rows if x['category']=='Процессор')
    assert _manual_cpu['alternatives'] and _manual_cpu['name'].startswith('AMD '), _manual_cpu
    assert any(x['Проверка']=='M.2 Key E' and x['Статус']=='OK' and 'переходник' in str(x['Факт QX']).lower() for x in _manual_checks), _manual_checks
    _adapter = next(x for x in _manual_rows if x['category']=='Переходник M.2 Key E')
    assert float(_adapter['price']) == 0.0 and _adapter['constraints'].get('provides_m2_key_e') == 1, _adapter

    _auto_same = Requirements(product_type='system', quantity=20, ram_type='DDR4', ram_gb_min=8,
                              ram_mhz_min=3200, cpu_cores_min=4, cpu_threads_min=8, cpu_freq_ghz_min=3.6,
                              cpu_l3_mb_min=4, cpu_tdp_w_max=90, need_integrated_graphics=True,
                              m2_key_m_min=1, m2_key_e_min=1, psu_w_min=400)
    _auto_plat,_auto_warn = resolve_platform(_auto_same)
    assert _auto_plat == 'amd', (_auto_plat, _auto_warn)

    # v2.7 regression: one corporate Word table can contain several goods, with
    # merged item cells and one row per characteristic. It must become four positions,
    # not one generic "ТИП 1", and optional "Да или нет" must remain unconstrained.
    from docx import Document as _GroupDoc
    from io import BytesIO as _GroupBIO
    _gd=_GroupDoc(); _gt=_gd.add_table(rows=1, cols=7)
    _hdr=['№ п/п','Наименование','Ед. Измер','Кол-во','Наименование показателя','Содержание (значение) показателя','Инструкция участнику закупки по формированию предложения']
    for _i,_v in enumerate(_hdr): _gt.rows[0].cells[_i].text=_v
    _grow=[
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Количество ядер процессора, шт','≥ 6','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Количество потоков процессора, шт','≥ 12','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Базовая тактовая частота производительных ядер процессора, Гигагерц','≥ 2,5','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Максимальная тактовая частота процессора, Гигагерц','≥ 4,4','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Объем кэш памяти процессора, Мегабайт','≥ 18','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Общий объем установленной оперативной памяти, Гигабайт','≥ 16','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Тип оперативной памяти','DDR5','Неизменный показатель'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Количество слотов оперативной памяти, шт','≥ 2','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Объем SSD накопителя тип 1, Гигабайт','≥ 240','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Объем HDD накопителя тип 2, Терабайт','≥ 1','Конкретное значение'),
        ('1','Системный блок\nОКПД2: 26.20.15.120','Шт.','54','Наличие встроенного картридера','Да или нет','Конкретное значение'),
        ('2','Монитор, подключаемый к компьютеру\nОКПД2: 26.20.17.110','шт','54','Размер диагонали, Дюйм (25,4 мм)','≥ 23.8','Конкретное значение'),
        ('2','Монитор, подключаемый к компьютеру\nОКПД2: 26.20.17.110','шт','54','Максимальная частота обновления (смена кадров), Герц','≥ 75','Конкретное значение'),
        ('2','Монитор, подключаемый к компьютеру\nОКПД2: 26.20.17.110','шт','54','Время отклика, мс','Не более 6','Конкретное значение'),
        ('2','Монитор, подключаемый к компьютеру\nОКПД2: 26.20.17.110','шт','54','Потребляемая мощность, Ватт','Не более 40','Конкретное значение'),
        ('3','Клавиатура компьютерная\nОКПД2: 26.20.16.110','шт','54','Интерфейс клавиатуры','USB','Неизменный показатель'),
        ('3','Клавиатура компьютерная\nОКПД2: 26.20.16.110','шт','54','Количество клавиш клавиатуры, шт','≥ 104','Конкретное значение'),
        ('3','Клавиатура компьютерная\nОКПД2: 26.20.16.110','шт','54','Передача данных клавиатуры','проводная','Неизменный показатель'),
        ('4','Мышь компьютерная\nОКПД2: 26.20.16.170','шт','54','Интерфейс манипулятора мышь','USB','Неизменный показатель'),
        ('4','Мышь компьютерная\nОКПД2: 26.20.16.170','шт','54','Количество кнопок манипулятора мышь','≥ 3 штука','Конкретное значение'),
        ('4','Мышь компьютерная\nОКПД2: 26.20.16.170','шт','54','Передача данных манипулятора мышь','проводная','Неизменный показатель'),
    ]
    for _vals in _grow:
        _cells=_gt.add_row().cells
        for _i,_v in enumerate(_vals): _cells[_i].text=_v
    _gb=_GroupBIO(); _gd.save(_gb)
    _greqs,_=parse_tz('grouped.docx',_gb.getvalue())
    assert len(_greqs)==4, [(x.title,x.product_type) for x in _greqs]
    _sys,_mon,_kb,_ms=_greqs
    assert _sys.product_type=='system' and _sys.quantity==54 and _sys.cpu_l3_mb_min==18
    assert _sys.ram_gb_min==16 and _sys.ram_type=='DDR5' and _sys.ram_slots_min==2
    assert _sys.ssd_gb_min==240 and _sys.hdd_gb_min==1000 and _sys.need_cardreader is False
    assert _mon.product_type=='monitor' and _mon.monitor_inches==23.8 and _mon.monitor_refresh_hz_min==75
    assert _mon.monitor_response_ms_max==6 and _mon.monitor_response_ms_max_inclusive is True and _mon.monitor_power_w_max==40
    assert _kb.product_type=='keyboard' and _kb.keyboard_usb and _kb.keyboard_wired and _kb.keyboard_keys_min==104
    assert _ms.product_type=='mouse' and _ms.mouse_usb and _ms.mouse_wired and _ms.mouse_buttons_min==3

    # v2.9.2 UI-state regression: native theme switching must carry the entire parsed/calculated
    # project plus current workspace to the other pre-warmed backend. A wrong-theme backend
    # must never consume the handoff file.
    from core.theme_handoff import save_handoff as _save_handoff, load_handoff as _load_handoff
    from tempfile import TemporaryDirectory as _UITemp
    from pathlib import Path as _UIPath
    with _UITemp() as _td:
        _hp=_UIPath(_td)/'theme_handoff.json'
        _pb=dump_project(reqs=[{'title':'Theme QA','quantity':2}],plans=[{'req':{'title':'Theme QA','quantity':2},'rows':[{'name':'Manual price','price':987.65}]}],source_filename='theme.docx')
        _save_handoff(_hp,target_theme='light',project_bytes=_pb,ui_state={'main_nav':'🧮 Расчёт','calc_workspace_stage':'🧩 Сборка и цены','sidebar_width':350})
        assert _load_handoff(_hp,current_theme='dark',consume=False) is None and _hp.exists()
        _hr=_load_handoff(_hp,current_theme='light',consume=True)
        assert _hr and not _hp.exists()
        _ho=load_project(_hr['project_bytes'])
        assert _ho['source_filename']=='theme.docx' and _ho['reqs'][0]['title']=='Theme QA'
        assert _ho['plans'][0]['rows'][0]['price']==987.65
        assert _hr['ui_state']['calc_workspace_stage']=='🧩 Сборка и цены'

    # Static UI guardrails: do not reintroduce the StreamlitWidgetAlreadyInstantiatedError
    # by assigning the stage widget after it has been rendered, and keep the dual-theme backend.
    _app_src=(_UIPath(__file__).resolve().parent/'app.py').read_text(encoding='utf-8')
    _launcher_src=(_UIPath(__file__).resolve().parent/'desktop_launcher.py').read_text(encoding='utf-8')
    assert "st.session_state.calc_workspace_stage='🧩 Сборка и цены'" not in _app_src
    assert "_calc_workspace_stage_pending" in _app_src and "save_theme_handoff" in _app_src and "load_theme_handoff" in _app_src
    assert "ports = {'dark': _free_port(), 'light': _free_port()}" in _launcher_src
    assert "QX_FORCED_THEME" in _launcher_src and "qx_theme_switch" in _launcher_src

    # v2.9.3 WebView form-control regression: values must stay explicitly visible even
    # when Chromium does not inherit Streamlit's native input text color until focus.
    for _selector_guard in (
        '[data-testid="stNumberInput"] input',
        '[data-testid="stTextInput"] input',
        '[data-baseweb="select"] [role="combobox"]',
        '-webkit-text-fill-color:var(--qx-input-text)!important',
        "color-scheme:{_PAL['scheme']}!important",
    ):
        assert _selector_guard in _app_src, _selector_guard

    def _lum(_hex):
        _h=_hex.lstrip('#'); _rgb=[int(_h[i:i+2],16)/255 for i in (0,2,4)]
        _rgb=[v/12.92 if v<=0.04045 else ((v+0.055)/1.055)**2.4 for v in _rgb]
        return .2126*_rgb[0]+.7152*_rgb[1]+.0722*_rgb[2]
    def _contrast(_a,_b):
        _x,_y=_lum(_a),_lum(_b); _hi,_lo=max(_x,_y),min(_x,_y)
        return (_hi+.05)/(_lo+.05)
    assert _contrast('#f8fafc','#18202d') >= 7.0
    assert _contrast('#111827','#ffffff') >= 7.0
    assert _contrast('#aab4c3','#18202d') >= 4.5
    assert _contrast('#66758a','#ffffff') >= 4.5
    # Input outlines are UI boundaries, so require at least 3:1 against their field background.
    assert _contrast('#5f7085','#ffffff') >= 3.0
    assert _contrast('#7890ad','#151e2b') >= 3.0

    # v2.9.4 border regression: BaseWeb inputs must have a real border/outline, not just border-color.
    for _border_guard in (
        'border:1px solid var(--qx-input-border)!important',
        'outline:1px solid var(--qx-input-border)!important',
        "input_border='#5f7085'",
        "input_border='#7890ad'",
        'border:1.5px solid var(--qx-input-border)!important',
    ):
        assert _border_guard in _app_src, _border_guard

    # v2.9.5 fullscreen regression: always-visible explicit controls open a table-only
    # focus mode. Returning restores the exact workspace and scrolls back to the table anchor.
    from core.ui_state import CALC_STAGES as _CALC_STAGES, BUILD_STAGE as _BUILD_STAGE, normalize_calc_stage as _norm_stage, make_focus_return_state as _focus_state, merge_ui_state_for_handoff as _merge_ui
    assert _norm_stage('✅ Проверка заявки') == '✅ Проверка заявки'
    assert _norm_stage('broken-stage') == _CALC_STAGES[0]
    for _focus_guard in (
        '⛶ Таблица на весь экран',
        '_enter_fullscreen_calc(i,_stage)',
        '_enter_fullscreen_tender(_stage)',
        '_render_fullscreen_component_editor()',
        '_render_fullscreen_tender_editor()',
        '_focus_scroll_restore_pending',
        "_maybe_restore_focus_scroll('qx-stage-build-anchor')",
        "_maybe_restore_focus_scroll('qx-stage-verify-anchor')",
        'scrollIntoView',
        "['_focus_return_anchor'] = f'qx-build-table-{int(idx)}'",
        "['_focus_return_anchor'] = 'qx-verify-table'",
        "_maybe_restore_focus_scroll(f'qx-build-table-{i}')",
        "_maybe_restore_focus_scroll('qx-verify-table')",
    ):
        assert _focus_guard in _app_src, _focus_guard

    # Theme switching performance guardrail: both backends warm in parallel and readiness is
    # cached, so repeated toggles do not block on a fresh Streamlit startup/probe.
    assert "ready_events = {'dark': threading.Event(), 'light': threading.Event()}" in _launcher_src
    assert "threading.Thread(target=_warm_backend" in _launcher_src

    # Coarse AST guard: never assign a literal session_state key after a literal-key widget
    # with the same key has already been instantiated earlier in app.py. This is exactly the
    # class of bug that previously broke workspace switching and full-screen return.
    import ast as _ast
    _tree=_ast.parse(_app_src)
    _widget_lines={}
    for _node in _ast.walk(_tree):
        if isinstance(_node,_ast.Call):
            for _kw in _node.keywords:
                if _kw.arg=='key' and isinstance(_kw.value,_ast.Constant) and isinstance(_kw.value.value,str):
                    _widget_lines.setdefault(_kw.value.value,[]).append(_node.lineno)
    _state_assignments=[]
    for _node in _ast.walk(_tree):
        if isinstance(_node,(_ast.Assign,_ast.AnnAssign,_ast.AugAssign)):
            _targets=_node.targets if isinstance(_node,_ast.Assign) else [_node.target]
            for _tg in _targets:
                if isinstance(_tg,_ast.Subscript) and isinstance(_tg.value,_ast.Attribute) and isinstance(_tg.value.value,_ast.Name) and _tg.value.value.id=='st' and _tg.value.attr=='session_state':
                    _sl=_tg.slice
                    if isinstance(_sl,_ast.Constant) and isinstance(_sl.value,str):
                        _state_assignments.append((_sl.value,_node.lineno))
                elif isinstance(_tg,_ast.Attribute) and isinstance(_tg.value,_ast.Attribute) and isinstance(_tg.value.value,_ast.Name) and _tg.value.value.id=='st' and _tg.value.attr=='session_state':
                    _state_assignments.append((_tg.attr,_node.lineno))
    _late=[]
    for _key,_line in _state_assignments:
        if _key in _widget_lines and any(_wl < _line for _wl in _widget_lines[_key]):
            _late.append((_key,_line,_widget_lines[_key]))
    assert not _late, _late

    print('QX Selector self-test: OK')


if __name__ == '__main__':
    main()

# v1.6 native file saving: test writer without opening GUI dialogs.
from core.native_files import write_bytes
from tempfile import TemporaryDirectory
from pathlib import Path as _Path
with TemporaryDirectory() as _td:
    _p = _Path(_td) / 'test_save.qxproj'
    _saved = write_bytes(_p, b'QX-DESKTOP-TEST')
    assert _Path(_saved).read_bytes() == b'QX-DESKTOP-TEST'

