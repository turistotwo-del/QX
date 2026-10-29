"""Repeatable QX Modern API smoke/regression suite. Run: python modern/test_modern.py.
The tests avoid writing real catalog rows, prices, or project autosave data.
"""
import base64
from dataclasses import asdict
from io import BytesIO
from pathlib import Path
import sys
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from docx import Document
from fastapi.testclient import TestClient
from core import storage
from core.tz_parser import Requirements
from modern.server import app


class ModernIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app, raise_server_exceptions=False)
        cls.prices = storage.load_prices()
        cls.settings = storage.load_settings()

    def request_ok(self, method, url, **kwargs):
        resp = self.client.request(method, url, **kwargs)
        self.assertEqual(200, resp.status_code, f'{url}: {resp.text[:800]}')
        return resp

    def sample_plan(self, platform='amd'):
        req = asdict(Requirements(title='Тестовый системный блок', quantity=3,
                                  platform=platform, cpu_cores_min=6,
                                  ram_gb_min=16, ssd_gb_min=480, m2_key_e_min=1))
        body = {'reqs': [req], 'prices': self.prices, 'settings': self.settings}
        return req, self.request_ok('POST', '/api/select', json=body).json()['plans'][0]

    def test_health_boot_static(self):
        self.assertTrue(self.request_ok('GET', '/api/health').json()['ok'])
        boot = self.request_ok('GET', '/api/boot').json()
        self.assertEqual('2.9.5', boot['engine_version'])
        self.assertIn('equipment_profile', boot['catalogs'])
        self.assertIn('app-shell', self.request_ok('GET', '/style.css').text)
        self.assertIn('QX Подборщик — Modern', self.request_ok('GET', '/').text)
        self.assertIn('QX Modern', self.request_ok('GET', '/app.js').text)

    def test_platform_hard_choice_and_adapter(self):
        for platform in ('amd', 'intel'):
            req, plan = self.sample_plan(platform)
            self.assertEqual(platform, plan['req']['platform'])
            self.assertGreater(len(plan['rows']), 0)
            self.assertIsInstance(plan['checks'], list)
            if platform == 'amd':
                adapter = [x for x in plan['rows'] if 'key e' in (x['name'] or '').lower()]
                self.assertTrue(adapter, 'AMD must include the standard Key E adapter')
                self.assertEqual(0, adapter[0]['price'])

    def test_currency(self):
        calc = self.request_ok('POST', '/api/price', json={
            'value': '$100', 'rate': 90, 'buffer': 5}).json()
        self.assertAlmostEqual(9500, calc['price'], places=2)
        self.assertEqual('USD', calc['currency'])
        calc = self.request_ok('POST', '/api/price', json={
            'value': '2500', 'rate': 90, 'buffer': 5}).json()
        self.assertAlmostEqual(2500, calc['price'], places=2)

    def test_manual_price_and_soft_reselect(self):
        req, plan = self.sample_plan()
        p = plan['rows'][0]
        p['price'] = 3456
        p['price_input'] = '$36.38'
        p['_locked'] = True
        plan['rows'].append({'category': 'Своя позиция', 'name': 'Моя доставка',
                             'qty': 2, 'price': 400, '_manual': True})
        body = {'reqs': [req], 'prices': self.prices, 'settings': self.settings,
                'previous': [plan], 'hard': False}
        selected = self.request_ok('POST', '/api/select', json=body).json()['plans'][0]
        self.assertEqual(3456, selected['rows'][0]['price'])
        self.assertEqual('$36.38', selected['rows'][0]['price_input'])
        self.assertTrue(any(x.get('_manual') for x in selected['rows']))
        self.assertEqual(2, selected['rows'][-1]['qty'])
        body['hard'] = True
        selected = self.request_ok('POST', '/api/select', json=body).json()['plans'][0]
        self.assertFalse(any(x.get('_manual') for x in selected['rows']))

    def test_project_backward_compatible(self):
        req, plan = self.sample_plan()
        body = {'name': 'Тестовый расчёт', 'reqs': [req], 'plans': [plan],
                'settings': self.settings, 'prices': self.prices}
        raw = self.request_ok('POST', '/api/project/save', json=body).content
        project = self.request_ok('POST', '/api/project/load', content=raw).json()
        self.assertEqual('QX_SELECTOR_PROJECT', project['format'])
        self.assertEqual(2, project['schema'])
        self.assertEqual('Тестовый расчёт', project['project_meta']['name'])
        self.assertEqual(3, project['reqs'][0]['quantity'])
        self.assertEqual(len(plan['rows']), len(project['plans'][0]['rows']))

    def test_exports_arsenal_and_docx_parse(self):
        req, plan = self.sample_plan()
        b = self.request_ok('POST', '/api/export/excel', json={'name': 'QX тест', 'plan': plan}).content
        self.assertTrue(b.startswith(b'PK'), 'XLSX must be a valid ZIP container')
        b = self.request_ok('POST', '/api/export/offer', json={
            'plans': [plan], 'percent': 20, 'warranty': 12, 'validity': 30}).content
        self.assertTrue(b.startswith(b'PK'), 'DOCX must be a valid ZIP container')
        text = self.request_ok('POST', '/api/arsenal', json={'plans': [plan]}).json()
        self.assertTrue(text['text'])
        self.assertGreaterEqual(text['purchase'], 0)
        d = Document()
        t = d.add_table(rows=2, cols=4)
        for j, value in enumerate(('Наименование характеристики', 'Значение характеристики', 'Единица измерения', 'Инструкция по заполнению характеристики')):
            t.cell(0, j).text = value
        for j, value in enumerate(('Количество ядер процессора', '≥ 6', 'Штука', 'Участник закупки указывает в заявке конкретное значение характеристики')):
            t.cell(1, j).text = value
        io = BytesIO(); d.save(io)
        parsed = self.request_ok('POST', '/api/parse', content=io.getvalue(),
                 headers={'X-File-Name': 'test.docx'}).json()
        self.assertTrue(parsed['reqs'])
        self.assertTrue(parsed['source_docx_base64'])
        self.assertEqual(io.getvalue(), base64.b64decode(parsed['source_docx_base64']))
        self.assertTrue(parsed['characteristics'])
        filled = self.request_ok('POST', '/api/export/filled-tz', json={
            'reqs': parsed['reqs'], 'plans': [plan],
            'source_docx_base64': parsed['source_docx_base64'],
            'characteristics': parsed['characteristics'],
            'overrides': {}}).content
        self.assertTrue(filled.startswith(b'PK'))

    def test_arsenal_four_equal_quantities_are_59_sets_not_236(self):
        # 59 systems + 59 monitors + 59 keyboards + 59 mice form 59 sets.
        # Prices here are a synthetic accounting fixture (NOT real catalog prices).
        items=[('Системный блок', 40000), ('Монитор', 11000),
               ('Клавиатура', 1000), ('Мышь', 1155.64)]
        plans=[{'req':{'title':name,'quantity':59},'rows':[{'qty':1,'price':price}]}
               for name,price in items]
        result=self.request_ok('POST','/api/arsenal',json={
            'plans':plans,'fields':{'delivery':40000}}).json()
        self.assertAlmostEqual(3136182.76,result['purchase'],places=2)
        self.assertEqual(59,result['summary']['set_quantity'])
        self.assertAlmostEqual(53155.64,result['summary']['set_unit_cost'],places=2)
        self.assertEqual(4,len(result['summary']['positions']))
        self.assertIn('53 155,64 за комплект',result['text'])
        self.assertIn('Количество комплектов: 59',result['text'])
        self.assertIn('3 136 182,76 общ.',result['text'])
        self.assertNotIn('13 288,91 за ед.',result['text'])
        self.assertAlmostEqual(3960228.45,result['exit_points'][0][1],places=2)

    def test_arsenal_different_quantities_no_fake_set_unit_price(self):
        plans=[{'req':{'title':'Системный блок','quantity':59},'rows':[{'qty':1,'price':30000}]},
               {'req':{'title':'Монитор','quantity':58},'rows':[{'qty':1,'price':9000}]}]
        result=self.request_ok('POST','/api/arsenal',json={'plans':plans}).json()
        self.assertIsNone(result['summary']['set_quantity'])
        self.assertIsNone(result['summary']['set_unit_cost'])
        self.assertFalse(result['summary']['is_set'])
        self.assertAlmostEqual(59*30000+58*9000,result['purchase'])
        self.assertIn('разные количества товаров',result['text'])
        self.assertIn('Системный блок: 59 шт.',result['text'])
        self.assertIn('Монитор: 58 шт.',result['text'])
        self.assertNotIn('за комплект',result['text'])

    def test_arsenal_single_item_and_bom_quantity(self):
        plan={'req':{'title':'Монитор','quantity':59},'rows':[{'qty':2,'price':1500},
             {'qty':1,'price':8000}]}
        result=self.request_ok('POST','/api/arsenal',json={'plans':[plan]}).json()
        self.assertAlmostEqual(11000,result['summary']['set_unit_cost'])
        self.assertAlmostEqual(649000,result['purchase'])
        self.assertIn('11 000 за ед.',result['text'])
        self.assertNotIn('за комплект',result['text'])

    def test_input_validation(self):
        result = self.client.post('/api/parse', content=b'data', headers={'X-File-Name':'bad.exe'})
        self.assertEqual(400, result.status_code)
        result = self.client.get('/api/catalog/not-real')
        self.assertEqual(400, result.status_code)
        result = self.client.post('/api/project/load', content=b'not json')
        self.assertEqual(400, result.status_code)

    def test_project_download_mime_and_zero_quantity(self):
        req, plan = self.sample_plan()
        data = {'name':'Тест загрузки проекта','reqs':[req], 'plans':[plan], 'prices':self.prices, 'settings':self.settings}
        result=self.request_ok('POST','/api/project/save',json=data)
        self.assertEqual('application/octet-stream',result.headers['content-type'])
        self.assertEqual('QX_SELECTOR_PROJECT',__import__('json').loads(result.content)['format'])
        plan['req']['quantity']=0
        result=self.client.post('/api/export/offer',json={'plans':[plan],'percent':20})
        self.assertEqual(400,result.status_code)
        self.assertIn('количество',result.text.lower())

    def test_parser_other_formats(self):
        for ext in ('txt','md'):
            payload=('Системный блок, количество 3 штуки.\nКоличество ядер процессора не менее 6.\nОбъём оперативной памяти 16 ГБ.').encode('utf-8')
            parsed=self.request_ok('POST','/api/parse',content=payload,headers={'X-File-Name':'test.'+ext}).json()
            self.assertTrue(parsed['reqs'])
            self.assertTrue(parsed['source_text'])

    def test_factual_ports_and_key_e_adapter(self):
        req, plan = self.sample_plan('amd')
        # Use the actual motherboard and adapter from the calculated plan.
        rows = [
            ('Количество слотов M.2 Key E','≥ 1','1'),
            ('Количество портов HDMI','≥ 1','1'),
            ('Общее количество установленных портов DisplayPort','≥ 1','1'),
            ('Количество сетевых интерфейсов 8P8C (RJ-45)','≥ 1','1'),
        ]
        ch=[{'id':str(i),'plan_index':0,'name':a,'original_value':b,'unit':'Штука','editable':True}
            for i,(a,b,_) in enumerate(rows)]
        actual=self.request_ok('POST','/api/characteristics',json={
            'characteristics':ch,'reqs':[req],'plans':[plan],'overrides':{}}).json()['rows']
        self.assertEqual([x[2] for x in rows],[r['value'] for r in actual])
        self.assertTrue(all(r['status']=='OK' for r in actual))

    def test_custom_chassis_prefers_registry_wifi(self):
        req=asdict(Requirements(title='Другой корпус',product_type='system',platform='amd',
            quantity=2,system_case_mode='custom',custom_case_name='Мой корпус',
            custom_case_height_mm=350,custom_case_length_mm=330,custom_case_width_mm=185))
        p=self.request_ok('POST','/api/select',json={
            'reqs':[req],'prices':self.prices,'settings':self.settings}).json()['plans'][0]
        self.assertTrue(any(r['category']=='Корпус' and r['name']=='Мой корпус' for r in p['rows']))
        self.assertTrue(any('Wi-Fi/Bluetooth QX' in r['name'] for r in p['rows']))
        self.assertFalse(any(r['category']=='Картридер' for r in p['rows']))

    def test_case_with_prohibited_wifi_uses_cardreader(self):
        req=asdict(Requirements(title='Без беспроводной связи',product_type='system',
            platform='amd',quantity=2,forbid_wireless=True))
        p=self.request_ok('POST','/api/select',json={
            'reqs':[req],'prices':self.prices,'settings':self.settings}).json()['plans'][0]
        self.assertFalse(p['registry'].get('registry_wifi'))

    def test_partial_arsenal_fields_and_kp_header(self):
        req,plan=self.sample_plan('amd')
        self.assertTrue(self.request_ok('POST','/api/arsenal',json={'plans':[plan],'fields':{}}).json()['text'])
        from core.commercial_offer import build_commercial_offer
        blob=self.request_ok('POST','/api/export/offer',json={
            'plans':[plan],'percent':20,'warranty':12,'validity':30}).content
        doc=Document(BytesIO(blob))
        self.assertGreaterEqual(len(doc.tables),2)
        self.assertIn('БРОКТЕХНОЛОДЖИ',doc.tables[0].cell(0,1).text)
        self.assertIn('Реестровый номер',doc.tables[1].cell(1,1).text)
        self.assertNotIn('БРОКТЕХНОЛОДЖИ',doc.tables[0].cell(0,0).text)

    def test_item_classification_monoblock_vs_monitor(self):
        from core.tz_parser import parse_one
        specs=[('Диагональ экрана','≥ 27','Дюйм'),('Количество ядер процессора','≥ 6','Штука'),
               ('Объем установленной оперативной памяти','≥ 16','Гигабайт'),('Интерфейс SSD','NVMe','')]
        common='\n'.join(' | '.join((a,b,u,'Участник указывает конкретное значение')) for a,b,u in specs)
        scenarios=[('Наименование товара: Моноблок 27 дюймов\nВнутри корпуса установлен системный блок\n'+common,'aio27'),
                   (common,'aio27'),
                   ('Наименование товара: Монитор\nДиагональ экрана | ≥ 27 | Дюйм\nРазрешение экрана | 1920x1080 | Пиксель','monitor'),
                   ('Наименование товара: Системный блок\nКоличество ядер процессора | ≥ 6 | Штука','system')]
        for source,kind in scenarios:
            with self.subTest(kind=kind,source=source[:55]):
                result=parse_one(source)
                self.assertEqual(kind,result.product_type)

    def test_monoblock_docx_uses_aio_plan_not_system_chassis(self):
        from core.tz_parser import parse_tz
        d=Document();d.add_paragraph('Наименование товара: Моноблок')
        t=d.add_table(rows=1, cols=4)
        for name,value,unit in [('Диагональ экрана','≥ 27','Дюйм'),
                                ('Количество ядер процессора','≥ 6','Штука'),
                                ('Объем установленной оперативной памяти','≥ 16','Гигабайт'),
                                ('Интерфейс SSD','NVMe','')]:
            cells=t.add_row().cells
            for cell,txt in zip(cells,(name,value,unit,'Участник закупки указывает конкретное значение')):
                cell.text=txt
        payload=BytesIO();d.save(payload)
        parsed=self.request_ok('POST','/api/parse',content=payload.getvalue(),
                   headers={'X-File-Name':'monoblock.docx'}).json()
        req=parsed['reqs'][0]
        self.assertEqual('aio27',req['product_type'])
        req['quantity']=10
        plan=self.request_ok('POST','/api/select',json={'reqs':[req],
                         'settings':self.settings,'prices':self.prices}).json()['plans'][0]
        self.assertEqual(10,plan['req']['quantity'])
        self.assertTrue(any(row['category']=='Моноблочная платформа' for row in plan['rows']))
        self.assertFalse(any(row['category']=='Корпус' for row in plan['rows']))


if __name__ == '__main__':
    unittest.main(verbosity=2)
