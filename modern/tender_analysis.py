"""Conservative document evidence extraction for procurement review.

This is deterministic local extraction, NOT an LLM. It never fabricates clauses,
contract deadlines or registry status. Optional local Ollama can be used with user consent.
"""
from __future__ import annotations
from io import BytesIO
import re
from pathlib import Path
from typing import Iterable
from html import unescape

REGISTRY_BIOS = re.compile(r'(?i)(?:реестр\w*|минцифр\w*|ЕРРП)[^|;.!?]{0,90}(?:BIOS|UEFI|БИОС)|(?:BIOS|UEFI|БИОС)[^|;.!?]{0,90}(?:реестр\w*|минцифр\w*|ЕРРП)')
RUSSIAN_BIOS = re.compile(r'(?i)(?:российск\w*|отечественн\w*)[^|;.!?]{0,90}(?:BIOS|UEFI|БИОС)|(?:BIOS|UEFI|БИОС)[^|;.!?]{0,90}(?:российск\w*|отечественн\w*)')
BIOS_ANY = re.compile(r'(?i)\b(?:BIOS|UEFI|БИОС)\b')
NEGATION = re.compile(r'(?i)\b(?:не\s+требует(?:ся)?|не\s+предъявл\w*|отсутствует|не\s+нуж\w*|не\s+обязател\w*|без\s+требован\w*)\b')
KEYWORDS = {
 'Срок поставки': re.compile(r'(?i)(?:срок(?:и|а|ов)?\s+(?:поставки|исполнени\w*)|постав\w*\s+в\s+течение|не\s+позднее|до\s+\d{1,2}\.\d{1,2}\.\d{4})'),
 'Оплата': re.compile(r'(?i)(?:оплат\w*|расч[её]т\w*\s+с\s+поставщиком)'),
 'Место поставки': re.compile(r'(?i)(?:адрес\s+поставки|место\s+поставки|место\s+доставки)'),
 'Гарантия': re.compile(r'(?i)(?:гарантийн\w*\s+(?:срок|обязательств\w*)|гарантия\s+на\s+(?:товар|оборудование))'),
 'Обеспечение': re.compile(r'(?i)(?:обеспечение\s+(?:исполнения|контракта|заявки))'),
 'Цена контракта': re.compile(r'(?i)(?:цен[аы]\s+контракт\w*|начальн\w*\s+\(?максимальн\w*\)?\s+цен\w*)'),
}


MAX_SCAN_SEGMENTS = 50000
MAX_PDF_PAGES = 250


def _prioritize_segments(segments: list[dict], limit: int) -> list[dict]:
    """Sample the whole document without discarding important late clauses.

    Keep the early and final parts plus the most relevant source excerpts from
    the middle. Every retained excerpt keeps its ORIGINAL location. This is a
    bounded preview, not a claim that unselected context was legally analysed.
    """
    if len(segments) <= limit:
        return segments
    if limit < 4:
        return segments[-limit:]
    head_size = max(1, limit // 4)
    tail_size = max(1, limit // 4)
    indices = set(range(head_size)) | set(range(len(segments) - tail_size, len(segments)))
    capacity = limit - len(indices)

    def importance(text: str) -> int:
        # Give the exact BIOS/UEFI mention highest precedence, even if it is
        # near the end of a large table or after thousands of contract lines.
        if BIOS_ANY.search(text):
            return 100
        if any(p.search(text) for p in KEYWORDS.values()):
            return 15
        if re.search(r'(?i)\b(?:НМЦК|реестров\w*|контракт\w*|оплат\w*)\b', text):
            return 3
        return 0

    matching = []
    for idx in range(head_size, len(segments) - tail_size):
        score = importance(segments[idx]['text'])
        if score:
            matching.append((score, idx))
    # Spread equal-priority results through the document, so all findings do
    # not come from the opening pages of an enormous document.
    scores = sorted({x[0] for x in matching}, reverse=True)
    for score in scores:
        matches = [idx for sc, idx in matching if sc == score]
        slots = min(capacity, len(matches))
        if slots <= 0:
            break
        chosen = matches if slots == len(matches) else [matches[min(len(matches)-1, int((i + .5) * len(matches) / slots))] for i in range(slots)]
        indices.update(chosen)
        capacity = limit - len(indices)
    return [segments[idx] for idx in sorted(indices)]


def document_segments(filename: str, data: bytes, limit=3500):
    """Read the entire supported source range, then keep evidence-rich excerpts.

    A hard scan cap generates an explicit error, never a misleading 'not found'.
    """
    ext = Path(filename).suffix.lower()
    segments: list[dict] = []

    def push(s, loc):
        norm = re.sub(r'\s+', ' ', str(s or '')).strip()
        if not norm:
            return
        if len(segments) >= MAX_SCAN_SEGMENTS:
            raise ValueError('Документ слишком большой для полного анализа; требуется разделить файл или проверить вручную')
        segments.append({'text': norm[:9000], 'location': loc})

    if ext == '.docx':
        from docx import Document
        document = Document(BytesIO(data))
        for i, block in enumerate(document.element.body.iterchildren(), 1):
            if block.tag.endswith('}p'):
                text = ' '.join(el.text or '' for el in block.iter() if el.tag.endswith('}t'))
                push(text, f'абзац {i}')
            elif block.tag.endswith('}tbl'):
                for j, row in enumerate([r for r in block.iterchildren() if r.tag.endswith('}tr')], 1):
                    cells = [' '.join(el.text or '' for el in cell.iter() if el.tag.endswith('}t')) for cell in row.iterchildren() if cell.tag.endswith('}tc')]
                    push(' | '.join(cells), f'таблица (блок {i}), строка {j}')
        # Real tender templates sometimes place procurement metadata and
        # mandatory requirements in repeated section headers/footers. Keep
        # their evidence separate from body paragraphs so attribution is exact.
        visited_parts = set()
        for section_index, section in enumerate(document.sections, 1):
            for kind, container in (('верхний колонтитул', section.header),
                                    ('нижний колонтитул', section.footer)):
                # Multiple Word sections can link to the same header/footer.
                part_id = str(container.part.partname)
                if part_id in visited_parts:
                    continue
                visited_parts.add(part_id)
                for paragraph_index, paragraph in enumerate(container.paragraphs, 1):
                    push(paragraph.text, f'{kind}, раздел {section_index}, абзац {paragraph_index}')
                for table_index, table in enumerate(container.tables, 1):
                    for row_index, row in enumerate(table.rows, 1):
                        push(' | '.join(cell.text for cell in row.cells),
                             f'{kind}, раздел {section_index}, таблица {table_index}, строка {row_index}')
    elif ext == '.xml':
        from defusedxml import ElementTree as SafeET
        if len(data) > 20 * 1024 * 1024:
            raise ValueError('XML слишком большой (предел 20 МБ)')
        root = SafeET.fromstring(data)

        def short_tag(value: str) -> str:
            return value.rsplit('}', 1)[-1]

        # Use an actual element path, not an ambiguous sequential leaf number.
        # Sibling indices distinguish repeated contract amounts / lot details.
        stack = [(root, '/' + short_tag(root.tag) + '[1]')]
        while stack:
            element, xml_path = stack.pop()
            children = list(element)
            if not children:
                value = (element.text or '').strip()
                if value:
                    tag = short_tag(element.tag)
                    push(tag + ': ' + value, 'XML: ' + xml_path)
                continue
            counters: dict[str, int] = {}
            items = []
            for child in children:
                tag = short_tag(child.tag)
                counters[tag] = counters.get(tag, 0) + 1
                items.append((child, xml_path + '/' + tag + '[' + str(counters[tag]) + ']'))
            stack.extend(reversed(items))
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError(f'PDF содержит {len(reader.pages)} страниц; анализ более {MAX_PDF_PAGES} страниц требует отдельной обработки')
        for i, page in enumerate(reader.pages, 1):
            for j, line in enumerate((page.extract_text() or '').splitlines(), 1):
                push(line, f'страница {i}, строка {j}')
    elif ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in wb:
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    push(' | '.join(str(v) if v is not None else '' for v in row), f'лист «{sheet.title}», строка {i}')
        finally:
            wb.close()
    elif ext in ('.txt', '.md', '.csv'):
        raw = data.decode('utf-8-sig', errors='replace')
        for i, line in enumerate(raw.splitlines(), 1):
            push(line, f'строка {i}')
    elif ext in ('.doc', '.xls', '.rtf'):
        raise ValueError('Оригинал DOC/XLS/RTF сохранён; автоматический анализ этого формата недоступен. Проверьте вручную.')
    else:
        raise ValueError('Формат не поддерживается')

    # Add split-heading clauses on the FULL sequence first. If we first select
    # excerpts, disconnected sections could look adjacent to each other.
    completed = _complete_wrapped_clauses(segments, MAX_SCAN_SEGMENTS)
    return _prioritize_segments(completed, limit)



def _complete_wrapped_clauses(segments, limit):
    """Join a broken heading with the next line/paragraph without altering originals.

    Many contract PDFs split 'Срок поставки:' and the number into separate
    lines.  Only obvious colon-ended procurement/BIOS headings are joined;
    arbitrary adjacent clauses are never combined into a fake requirement.
    The evidence location explicitly records both original lines.
    """
    heading = re.compile(
        r'(?i)(?:срок\s+(?:поставки|исполнения)|место\s+поставки|адрес\s+поставки|'
        r'условия\s+оплаты|порядок\s+оплаты|гарантийный\s+срок|'
        r'требовани\w*\s+к\s+(?:BIOS|UEFI|БИОС)|(?:BIOS|UEFI|БИОС))\s*[:：]\s*$'
    )
    loc_patterns = (re.compile(r'^страница (\d+), строка (\d+)$'),
                    re.compile(r'^абзац (\d+)$'),
                    re.compile(r'^строка (\d+)$'))
    original = list(segments)
    added=[]
    for a,b in zip(original,original[1:]):
        if len(added)+len(original)>=limit:break
        first=a['text'].strip()
        second=b['text'].strip()
        if not first or not second or len(first)>180 or len(second)>600:
            continue
        if not heading.search(first):continue
        loc_a,loc_b=a['location'],b['location']
        adjacent=False
        for pat in loc_patterns:
            ma,mb=pat.match(loc_a),pat.match(loc_b)
            if ma and mb:
                ia,ib=tuple(map(int,ma.groups())),tuple(map(int,mb.groups()))
                adjacent=(len(ia)==len(ib)==1 and ib[0]==ia[0]+1 or
                          len(ia)==len(ib)==2 and ia[0]==ib[0] and ib[1]==ia[1]+1)
                break
        if not adjacent:continue
        # Refuse a new heading, which would make this a conjunction of unrelated facts.
        if re.match(r'(?i)^(?:\d+[.)]\s|(?:поставка|оплата|гарантия|реестровый BIOS)\s*:)',second):
            continue
        added.append({'text':first+' '+second,
                      'location':loc_a+' — '+loc_b+' (перенос строки)'})
    segments.extend(added)
    return segments

def extract_suggestions(segments):
    """Only strong, source-backed candidates. Conflicts/ambiguity are not applied."""
    supply=re.compile(r'(?i)(?:место|адрес)\s+(?:поставки|доставки)\s*[:—-]\s*(.{8,220})')
    term=re.compile(r'(?i)(?:срок\s+поставки|поставка\s+(?:товара|оборудования)?)\s*[:—-]?\s*(?:не\s+позднее\s+)?(?:в\s+течение\s+)?(\d{1,3})\s+(рабочих|календарных)\s+дн(?:ей|я|ь)\s+(?:с|со)\s+дат[ыи]\s+(?:заключени[яе]|подписани[яе])\s+контракт')
    guarantee=re.compile(r'(?i)(?:гарантийн(?:ый|ого)\s+(?:срок|период)|гарантия\s+на\s+(?:товар|оборудование))\s*[:—-]?\s*(.{3,170})')
    candidates=[]
    for item in segments:
        txt=item['text']
        m=term.search(txt)
        if m:
            number=int(m.group(1))
            if 0<number<=3650:
                candidates.extend([{'field':'due_days','value':number,'label':'Срок поставки (дни)','source':item},
                                  {'field':'due_unit','value':'work' if m.group(2).lower().startswith('раб') else 'calendar','label':'Тип дней','source':item}])
        m=supply.search(txt)
        if m:
            addr=m.group(1).split(' | ')[0].strip(' .;')[:190]
            if addr and not any(w in addr.lower() for w in ('значение характеристики','участник закупки')):
                candidates.append({'field':'supply_address','value':addr,'label':'Адрес поставки','source':item})
        for key,pat,label in (
            ('payment_terms',KEYWORDS['Оплата'],'Условия оплаты'),
            ('warranty',guarantee,'Гарантия'),
            ('security',KEYWORDS['Обеспечение'],'Обеспечение контракта')):
            if pat.search(txt) and len(txt)>=16:
                # No invented amount/date: use only the original clause, not our interpretation.
                candidates.append({'field':key,'value':txt[:250],'label':label,'source':item})
    accepted=[]
    for field in ('due_days','due_unit','supply_address','payment_terms','warranty','security'):
        choices=[c for c in candidates if c['field']==field]
        unique={str(c['value']).strip().lower() for c in choices}
        if len(unique)==1:
            c=choices[0]
            accepted.append({'field':field,'value':c['value'],'label':c['label'],
                             'evidence':{'file':c['source']['file'],'location':c['source']['location'],'text':c['source']['text'][:340]}})
        # Conflicting values across the documents -> no automatic proposal.
    return accepted


def analyze_documents(documents: Iterable[tuple[str,bytes]], max_evidence=3):
    documents = list(documents)
    segments=[]
    errors=[]
    for name,data in documents:
        try:
            # Do not make a compliance determination from a preview sample:
            # a single negation / incompatible clause can sit between thousands
            # of other BIOS references in a very large contract. Scan every
            # supported source segment; only the RETURNED evidence is limited.
            found=document_segments(name,data,limit=MAX_SCAN_SEGMENTS)
            if not found:
                errors.append({'file':name,
                               'error':'Нет доступного для анализа текста; возможен скан. Проверьте файл вручную.'})
                continue
            for s in found:
                segments.append({'file':name,**s})
        except Exception as exc:
            errors.append({'file':name,'error':str(exc)[:180]})
    evidence={}
    for key,pat in KEYWORDS.items():
        evidence[key]=[{**s, 'text':s['text'][:340]} for s in segments if pat.search(s['text'])][:max_evidence]
    candidates=[]
    russian=[]
    mentions=[]
    for s in segments:
        if REGISTRY_BIOS.search(s['text']): candidates.append(s)
        elif RUSSIAN_BIOS.search(s['text']): russian.append(s)
        elif BIOS_ANY.search(s['text']): mentions.append(s)
    positive=[s for s in candidates if not NEGATION.search(s['text'])]
    negative=[s for s in candidates if NEGATION.search(s['text'])]
    if errors:
        status='Требуется уточнение: не все документы удалось проанализировать'
    elif positive and negative: status='Требуется уточнение'
    elif positive: status='Найдено требование к реестровому BIOS'
    elif negative: status='Обнаружено отрицание требования — проверьте контекст'
    elif russian: status='Упомянут российский BIOS; реестровая запись не подтверждена'
    elif mentions: status='BIOS упоминается, реестровое требование не подтверждено'
    else: status='Требование не обнаружено в загруженных файлах'
    # Evidence must explain the STATUS, not just show the first few matching
    # mentions. Thousands of positive clauses could previously hide the one
    # contradictory/negative clause that triggered 'Требуется уточнение'.
    groups = (positive, negative, russian, mentions)
    prioritized = []
    seen_locations = set()
    for group in groups:
        if group:
            candidate=group[0]
            identity=(candidate['file'],candidate['location'])
            if identity not in seen_locations:
                prioritized.append(candidate)
                seen_locations.add(identity)
    for group in groups:
        for candidate in group[1:]:
            identity=(candidate['file'],candidate['location'])
            if identity not in seen_locations:
                prioritized.append(candidate)
                seen_locations.add(identity)
            if len(prioritized)>=max_evidence:
                break
        if len(prioritized)>=max_evidence:
            break
    bios=[{**s,'text':s['text'][:420]} for s in prioritized[:max_evidence]]
    summary=[]
    for k in ('Срок поставки','Место поставки','Оплата','Гарантия'):
        hits=evidence.get(k) or []
        summary.append({'label':k,'value':hits[0]['text'][:160] if hits else 'Не найдено','evidence':hits[:1]})
    summary.append({'label':'Реестровый BIOS','value':status,'evidence':bios[:1]})
    risk='Проверить подтверждение BIOS/UEFI выбранного оборудования' if positive or russian else 'Проверить сроки и существенные условия вручную'
    summary.append({'label':'Контроль','value':risk, 'evidence':bios[:1] if positive else []})
    return {'suggested_fields':extract_suggestions(segments), 'mode':'Локальное извлечение, без нейросети','summary':summary[:7], 'bios':{'status':status,'evidence':bios}, 'evidence':evidence, 'errors':errors, 'files_analyzed':len(documents) - len(errors),
            'disclaimer':'Найденные фразы не являются юридическим заключением. Для сроков и BIOS проверьте оригинальные пункты и действующую редакцию документации.'}
