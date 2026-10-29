"""Offline, source-linked Word review of one procurement case.

The report is evidence, not a compliance certificate. It does not make remote
requests and it never modifies the stored tender or its documents.
"""
from __future__ import annotations
from datetime import datetime
from io import BytesIO
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

FIELDS = (
    ('number', 'Номер закупки'), ('lot', 'Лот'), ('title', 'Предмет закупки'),
    ('customer', 'Заказчик'), ('law', 'Закон'), ('status', 'Статус (в QX)'),
    ('contract_number', 'Номер контракта'), ('signed_on', 'Дата подписания'),
    ('due_on', 'Срок поставки (дата)'), ('due_days', 'Срок поставки (дней)'),
    ('due_unit', 'Тип дней'), ('supply_address', 'Место поставки'),
    ('payment_terms', 'Условия оплаты'), ('warranty', 'Гарантия'),
    ('security', 'Обеспечение'),
)
FINANCE = (
    ('nmck', 'НМЦК'), ('bid_price', 'Цена предложения'),
    ('contract_price', 'Цена контракта'), ('purchase_cost', 'Себестоимость'),
    ('delivery_cost', 'Логистика'), ('extra_cost', 'Прочие расходы'),
    ('paid_amount', 'Получено оплат'),
)


def _value(value):
    if value is None or value == '':
        return 'Не заполнено'
    if isinstance(value, float):
        return f'{value:,.2f}'.replace(',', ' ')
    return str(value)[:1500]


def build_tender_report(tender: dict, analysis: dict, *, created_at=None) -> bytes:
    """Generate read-only DOCX using ONLY provided local case data/evidence."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)

    # Strict black-on-white office style; Word's default heading styles are blue.
    for heading_name, point_size in (('Heading 1', 13), ('Heading 2', 11)):
        heading_style = doc.styles[heading_name]
        heading_style.font.name = 'Times New Roman'
        heading_style.font.color.rgb = RGBColor(0, 0, 0)
        heading_style.font.size = Pt(point_size)
        heading_style.font.bold = True
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run=heading.add_run('ОТЧЁТ ПРОВЕРКИ ЗАКУПКИ')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = True
    date_line = (created_at or datetime.now().astimezone()).strftime('%d.%m.%Y %H:%M')
    p=doc.add_paragraph(f'QX Подборщик · локальный анализ документов · {date_line}')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(rows):
        tab=doc.add_table(rows=0, cols=2)
        tab.style='Table Grid'
        tab.alignment=WD_TABLE_ALIGNMENT.CENTER
        for label, value in rows:
            cells=tab.add_row().cells
            cells[0].text=str(label)
            cells[1].text=_value(value)
        return tab

    doc.add_heading('1. Карточка закупки', level=1)
    table([(label,tender.get(key)) for key,label in FIELDS])

    doc.add_heading('2. Финансовые сведения', level=1)
    table([(label, _value(tender.get(key)) + (' ₽' if tender.get(key) not in (None,'') else ''))
           for key,label in FINANCE])
    doc.add_paragraph('Финансовые суммы взяты из карточки QX и не подтверждаются автоматически документацией ЕИС.')

    docs=tender.get('documents') or []
    doc.add_heading('3. Использованные документы', level=1)
    if docs:
        for row in docs:
            doc.add_paragraph(str(row.get('filename') or 'Файл без названия')[:240], style='List Bullet')
    else:
        doc.add_paragraph('Документы не приложены. Анализ не выполнялся.')

    doc.add_heading('4. Реестровый BIOS / UEFI',level=1)
    bios=analysis.get('bios') or {}
    doc.add_paragraph('Результат: ' + str(bios.get('status') or 'Не проверено'))
    lines=bios.get('evidence') or []
    if lines:
        for e in lines:
            doc.add_paragraph(f"Источник: {e.get('file','?')} — {e.get('location','?')}")
            doc.add_paragraph(str(e.get('text') or '')[:1500])
    else:
        doc.add_paragraph('Подтверждающие выдержки не найдены.')

    doc.add_heading('5. Условия договора и выдержки',level=1)
    evidence=analysis.get('evidence') or {}
    if not evidence:
        doc.add_paragraph('Подтверждённых фрагментов нет.')
    for label, clauses in evidence.items():
        doc.add_heading(str(label), level=2)
        if not clauses:
            doc.add_paragraph('Не найдено в доступном для чтения тексте.')
        for e in clauses:
            doc.add_paragraph(f"{e.get('file','?')} · {e.get('location','?')}")
            doc.add_paragraph(str(e.get('text') or '')[:1500])

    doc.add_heading('6. Предложения по заполнению карточки', level=1)
    suggestions=analysis.get('suggested_fields') or []
    if suggestions:
        for item in suggestions:
            src=item.get('evidence') or {}
            doc.add_paragraph(f"{item.get('label',item.get('field',''))}: {_value(item.get('value'))}")
            doc.add_paragraph(f"Источник: {src.get('file','?')}, {src.get('location','?')}")
    else:
        doc.add_paragraph('Не найдено однозначных условий для автоматического предложения.')
    doc.add_paragraph('Предложения не применены к карточке и не заменяют проверки оригиналов.')

    doc.add_heading('7. Ограничения анализа',level=1)
    for err in analysis.get('errors') or []:
        doc.add_paragraph(f"НЕ ПРОВЕРЕН: {err.get('file','?')} — {err.get('error','нет текстового слоя')}",style='List Bullet')
    if not analysis.get('errors'):
        doc.add_paragraph('Необработанных файлов анализатор не сообщил. Это не подтверждает юридическую полноту проверки.')
    doc.add_paragraph(str(analysis.get('disclaimer') or 'Отчёт носит вспомогательный характер. Необходимо проверить оригиналы документов.'))
    doc.add_paragraph('QX не удостоверяет сведения ЕИС, реестровую запись BIOS, соответствие оборудования ТЗ или юридическую силу документации.')

    buf=BytesIO();doc.save(buf);return buf.getvalue()
