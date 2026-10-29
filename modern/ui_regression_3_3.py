"""QX Modern 3.3 real Chromium UI regression: actual monoblock procurement Word, live quantity + redraws."""
from pathlib import Path
from docx import Document
import tempfile
import threading,time
import uvicorn
from playwright.sync_api import sync_playwright
import requests
from modern.server import app

PORT=18759
server=uvicorn.Server(uvicorn.Config(app,host='127.0.0.1',port=PORT,log_level='error',access_log=False,log_config=None))
t=threading.Thread(target=server.run,daemon=True);t.start()
for _ in range(100):
 if server.started:break
 time.sleep(.1)
else:raise AssertionError('No server')
errors=[]
try:
 with sync_playwright() as p:
  browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
  page=browser.new_page(viewport={'width':1480,'height':860},accept_downloads=True)
  page.on('pageerror',lambda exc:errors.append(str(exc)))
  root=Path(__file__).resolve().parent/'web'
  def route(r):
   q=r.request
   url=q.url.replace('http://qx-ui.local',f'http://127.0.0.1:{PORT}')
   resp=requests.request(q.method,url,data=q.post_data_buffer,headers={k:v for k,v in q.headers.items() if k not in ('host','content-length')},timeout=25)
   r.fulfill(status=resp.status_code,body=resp.content,headers={'content-type':resp.headers.get('content-type','application/json')})
  page.route('**/api/**',route)
  page.set_content('<html><head><base href="http://qx-ui.local/"></head><body><div id="app"></div><input id="file-tz" type="file"><input id="file-project" type="file"><input id="file-catalog" type="file"><input id="file-supplier" type="file"></body></html>')
  page.add_style_tag(content=(root/'style.css').read_text())
  page.add_script_tag(content=(root/'app.js').read_text())
  page.wait_for_timeout(600)
  assert page.locator('h1').first.inner_text() == 'Рабочий стол'
  fixture=Path(tempfile.gettempdir())/'QX_modern_33_monoblock.docx'
  d=Document();d.add_paragraph('Наименование товара: Моноблок 27 дюймов')
  tbl=d.add_table(rows=1,cols=4)
  for row in [('Диагональ экрана','≥ 27','Дюйм'),('Количество ядер процессора','≥ 6','Штука'),('Объем установленной оперативной памяти','≥ 16','Гигабайт'),('Интерфейс SSD','NVMe','')]:
   cells=tbl.add_row().cells
   for cell,value in zip(cells,(*row,'Участник закупки указывает конкретное значение')):cell.text=value
  d.save(fixture)
  page.locator('#file-tz').set_input_files(str(fixture))
  page.locator('.loaded-source').wait_for(timeout=20000)
  page.wait_for_timeout(350)
  assert 'QX_modern_33_monoblock.docx' in page.locator('.loaded-source').inner_text()
  kind=page.locator('.product-classification strong').inner_text()
  assert 'Моноблок' in kind,kind
  assert page.locator('.requirement-group h2').filter(has_text='Экран и моноблочная платформа').count()==1
  assert page.locator('.requirement-group h2').filter(has_text='Характеристики монитора').count()==0
  assert page.locator('.quantity-warning').is_visible()
  quantity=page.locator('[data-req=quantity]')
  quantity.fill('1')
  assert not page.locator('.quantity-warning').is_visible()
  assert quantity.evaluate('(el)=>document.activeElement === el')
  quantity.fill('10')
  assert not page.locator('.quantity-warning').is_visible()
  assert quantity.input_value()=='10'
  quantity.press('Tab')
  assert page.locator('[data-req=quantity]').input_value()=='10'
  page.locator('[data-action=select]').first.click()
  page.locator('.price-table tbody tr').first.wait_for(timeout=20000)
  assert 'QX_modern_33_monoblock.docx' in page.locator('.source-strip').inner_text()
  page.locator('[data-build-quantity]').fill('4')
  assert page.locator('[data-build-quantity]').input_value()=='4'
  assert page.locator('[data-quantity-hint]').inner_text()=='Количество изделий в партии'
  page.locator('[data-build-quantity]').press('Tab')
  page.locator('[data-action=tab][data-tab=tz]').click()
  assert page.locator('[data-req=quantity]').input_value()=='4'
  page.screenshot(path=str(Path(tempfile.gettempdir())/'qx_3_3_monoblock_check.png'))
  page.locator('[data-action=tab][data-tab=settings]').click()
  page.locator('[data-action=theme][data-theme=light]').click()
  page.locator('[data-action=tab][data-tab=tz]').click()
  assert 'Моноблок' in page.locator('.product-classification strong').inner_text()
  assert not errors,errors
  print('CHROMIUM E2E PASS | monoblock classification, relevant groups, source filename, quantity live, tab persistence, build, light theme')
  browser.close()
finally:
 server.should_exit=True;t.join(timeout=6)
