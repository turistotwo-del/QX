from io import BytesIO
import json
from datetime import date, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _excel_value(value):
    """Convert arbitrary project values to something openpyxl can safely store."""
    if value is None:
        return ''
    if isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    if isinstance(value, (dict, list, tuple, set)):
        try:
            if isinstance(value, set):
                value = sorted(value, key=str)
            return json.dumps(value, ensure_ascii=False, indent=2, default=str)
        except Exception:
            return str(value)
    return str(value)


def export_xlsx(project_name, req, rows, registry, extras, checks=None):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Сборка'
    headers = ['Категория', 'Наименование', 'Кол-во на ед.', 'Цена за ед., ₽', 'Сумма на ед., ₽', 'Источник', 'Примечание', 'Ссылка']
    ws.append(headers)
    for r in rows:
        qty = float(r.get('qty', 1) or 0)
        price = float(r.get('price', 0) or 0)
        ws.append([
            _excel_value(r.get('category', '')),
            _excel_value(r.get('name', '')),
            qty,
            price,
            qty * price,
            _excel_value(r.get('source', '')),
            _excel_value(r.get('note', '')),
            _excel_value(r.get('url', '')),
        ])

    unit_cost = sum(float(r.get('qty', 1) or 0) * float(r.get('price', 0) or 0) for r in rows)
    quantity = int(req.get('quantity', 0) or 0)
    purchase = unit_cost * quantity
    delivery = float(extras.get('delivery', 0) or 0)
    bg1 = float(extras.get('bg1', 0) or 0)
    bg2 = float(extras.get('bg2', 0) or 0)
    markup = float(extras.get('markup', 0) or 0)
    total_cost = purchase + delivery + bg1 + bg2
    sale_total = total_cost * (1 + markup)

    ws2 = wb.create_sheet('Расчёт')
    for row in [
        ('Проект', project_name),
        ('Количество', quantity),
        ('Себестоимость 1 ед.', unit_cost),
        ('Стоимость закупки', purchase),
        ('Доставка', delivery),
        ('БГ №1', bg1),
        ('БГ №2', bg2),
        ('Общая себестоимость', total_cost),
        ('Наценка', markup),
        ('Цена предложения', sale_total),
        ('Цена предложения за 1 ед.', sale_total / quantity if quantity else 0),
        ('КТРУ', req.get('ktru_code')),
        ('НМЦК из ТЗ', req.get('nmck')),
    ]:
        ws2.append([_excel_value(v) for v in row])

    ws3 = wb.create_sheet('Реестр')
    if registry:
        for k, v in registry.items():
            ws3.append([_excel_value(k), _excel_value(v)])
    else:
        ws3.append(['Реестровая запись', 'Для этой позиции не подбирается'])

    ws4 = wb.create_sheet('Распознано из ТЗ')
    ws4.append(['Параметр', 'Значение'])
    for k, v in req.items():
        if k != 'raw_text':
            ws4.append([_excel_value(k), _excel_value(v)])

    if checks:
        ws5 = wb.create_sheet('Проверка QX')
        ws5.append(['Проверка', 'Требование', 'Факт QX', 'Статус'])
        for c in checks:
            ws5.append([
                _excel_value(c.get('Проверка')),
                _excel_value(c.get('Требование')),
                _excel_value(c.get('Факт QX')),
                _excel_value(c.get('Статус')),
            ])

    header_fill = PatternFill('solid', fgColor='111827')
    header_font = Font(color='FFFFFF', bold=True)
    thin = Side(style='thin', color='D1D5DB')
    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical='center')
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for col in range(1, sheet.max_column + 1):
            width = min(55, max(12, max((len(str(sheet.cell(r, col).value or '')) for r in range(1, min(sheet.max_row, 100) + 1)), default=12) + 2))
            sheet.column_dimensions[get_column_letter(col)].width = width

    for c in ws['D'] + ws['E']:
        c.number_format = '#,##0.00 ₽'

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
