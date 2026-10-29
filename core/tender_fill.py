from __future__ import annotations

from io import BytesIO
import re
from typing import Any, Dict, Iterable, List, Tuple

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P

from .cpu_catalog import cpu_details
from .defaults import QX_SPECS, CASE_SPECS, AIO_SPECS


def _norm(text: Any) -> str:
    return re.sub(r'\s+', ' ', str(text or '').replace('\xa0', ' ')).strip()


def _low(text: Any) -> str:
    return _norm(text).lower().replace('ё', 'е').replace('￾', '-')


def _find_header_columns(table: Table):
    for row_idx, row in enumerate(table.rows[:4]):
        cells = [_low(c.text) for c in row.cells]
        name_col = next((i for i, x in enumerate(cells) if 'наименование характерист' in x), None)
        value_col = next((i for i, x in enumerate(cells) if 'значение характерист' in x and 'инструкц' not in x), None)
        instruction_col = next((i for i, x in enumerate(cells) if 'инструкц' in x and 'заполн' in x), None)
        unit_col = next((i for i, x in enumerate(cells) if 'единиц' in x and 'измер' in x), None)
        if name_col is not None and value_col is not None and instruction_col is not None:
            return row_idx, name_col, value_col, unit_col, instruction_col
    return None


def _is_editable_instruction(instruction: str) -> bool:
    low = _low(instruction)
    if 'не может изменяться' in low:
        return False
    return ('участник закупки' in low and 'указывает' in low) or ('конкретное значение' in low)


def extract_characteristics_docx(data: bytes) -> List[dict]:
    """Extract characteristic rows with exact DOCX table coordinates.

    The coordinates are later used to fill only the value cell and leave all
    immutable characteristics untouched.
    """
    doc = Document(BytesIO(data))
    out: List[dict] = []
    table_index = -1
    current_position = 0
    saw_position_heading = False
    char_table_seq = 0

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            p = Paragraph(child, doc)
            txt = _norm(p.text)
            m = re.search(r'\bПозиция\s+(\d+)\b', txt, re.I)
            if m:
                current_position = max(0, int(m.group(1)) - 1)
                saw_position_heading = True
        elif isinstance(child, CT_Tbl):
            table_index += 1
            table = Table(child, doc)
            hdr = _find_header_columns(table)
            if not hdr:
                continue
            header_row, name_col, value_col, unit_col, instruction_col = hdr
            plan_index = current_position if saw_position_heading else char_table_seq
            char_table_seq += 1
            for row_idx in range(header_row + 1, len(table.rows)):
                row = table.rows[row_idx]
                if max(name_col, value_col, instruction_col, unit_col or 0) >= len(row.cells):
                    continue
                name = _norm(row.cells[name_col].text)
                value = _norm(row.cells[value_col].text)
                unit = _norm(row.cells[unit_col].text) if unit_col is not None and unit_col < len(row.cells) else ''
                instruction = _norm(row.cells[instruction_col].text)
                if not name:
                    continue
                editable = _is_editable_instruction(instruction)
                out.append({
                    'id': f't{table_index}r{row_idx}',
                    'table_index': table_index,
                    'row_index': row_idx,
                    'name_col': name_col,
                    'value_col': value_col,
                    'unit_col': unit_col,
                    'instruction_col': instruction_col,
                    'plan_index': plan_index,
                    'name': name,
                    'original_value': value,
                    'unit': unit,
                    'instruction': instruction,
                    'editable': editable,
                })
    return out


def _component(plan: dict, category: str):
    return next((r for r in (plan or {}).get('rows', []) if _low(r.get('category')) == _low(category)), None)


def _selected_cpu(plan: dict):
    row = _component(plan, 'Процессор')
    if not row:
        return None
    model = re.sub(r'\s+OEM\s*$', '', _norm(row.get('name')), flags=re.I)
    details = cpu_details([model])
    return details[0] if details else None


def _platform(plan: dict, req: dict):
    reg = (plan or {}).get('registry') or {}
    p = reg.get('platform') or (req or {}).get('platform')
    if p in QX_SPECS:
        return p
    mb = _component(plan, 'Материнская плата') or {}
    cons = dict(mb.get('constraints') or {})
    cp = cons.get('platform')
    if cp in ('amd', 'intel'):
        return cp
    name = _low(mb.get('name'))
    if 'b760' in name:
        return 'intel'
    if 'a520' in name:
        return 'amd'
    return None


def _board_facts(plan: dict, req: dict):
    """Facts of the motherboard that is actually present in the calculation.

    New projects carry the selected board specification inside the component row.
    This prevents a saved third-party board (e.g. Beshtau) from inheriting QX facts
    merely because both use the same AMD/Intel platform.
    """
    mb = _component(plan, 'Материнская плата') or {}
    cons = dict(mb.get('constraints') or {})
    if cons.get('kind') == 'motherboard':
        return cons
    platform = _platform(plan, req)
    return dict(QX_SPECS.get(platform or '', {}))


def _ram_facts(plan: dict):
    row = _component(plan, 'Оперативная память') or {}
    name = _norm(row.get('name'))
    cons = row.get('constraints') or {}
    ram_type = None
    m = re.search(r'\b(DDR[45])\b', name, re.I)
    if m:
        ram_type = m.group(1).upper()
    elif cons.get('ram_type'):
        ram_type = str(cons['ram_type']).upper()
    gb = None
    m = re.search(r'(\d+)\s*ГБ', name, re.I)
    if m:
        gb = int(m.group(1))
    elif cons.get('min_gb'):
        gb = int(cons['min_gb'])
    mhz = None
    m = re.search(r'(\d{3,5})\s*МГц', name, re.I)
    if m:
        mhz = int(m.group(1))
    elif cons.get('min_mhz'):
        mhz = int(cons['min_mhz'])
    modules = cons.get('modules')
    module_gb = cons.get('module_gb')
    m = re.search(r'(\d+)\s*[xх×]\s*(\d+)\s*ГБ', name, re.I)
    if m:
        modules, module_gb = int(m.group(1)), int(m.group(2))
    return {'type': ram_type, 'gb': gb, 'mhz': mhz, 'modules': modules, 'module_gb': module_gb}


def _ssd_facts(plan: dict):
    rows = [r for r in (plan or {}).get('rows', []) if _low(r.get('category')) == 'ssd']
    capacities = []
    interfaces = []
    for row in rows:
        name = _norm(row.get('name'))
        m = re.search(r'(\d+)\s*ГБ', name, re.I)
        if m:
            capacities.append(int(m.group(1)))
        cons = row.get('constraints') or {}
        iface = cons.get('interface') or ('NVMe' if re.search(r'\bnvme\b', name, re.I) else 'SATA' if re.search(r'\bsata\b', name, re.I) else '')
        if iface:
            interfaces.append(str(iface))
    return {
        'count': len(rows),
        'capacities': capacities,
        'total_gb': sum(capacities) if capacities else None,
        'capacity_gb': capacities[0] if capacities else None,
        'interface': interfaces[0] if interfaces else None,
    }


def _hdd_facts(plan: dict):
    rows = [r for r in (plan or {}).get('rows', []) if _low(r.get('category')) == 'hdd']
    capacities = []
    for row in rows:
        cons = row.get('constraints') or {}
        gb = cons.get('capacity_gb') or cons.get('min_gb')
        if gb:
            capacities.append(int(gb))
            continue
        name = _norm(row.get('name'))
        m = re.search(r'(\d+(?:[\.,]\d+)?)\s*ТБ', name, re.I)
        if m:
            capacities.append(int(round(float(m.group(1).replace(',', '.')) * 1000)))
            continue
        m = re.search(r'(\d+)\s*ГБ', name, re.I)
        if m:
            capacities.append(int(m.group(1)))
    return {
        'count': len(rows),
        'capacities': capacities,
        'total_gb': sum(capacities) if capacities else None,
        'capacity_gb': capacities[0] if capacities else None,
    }


def _drive_capacity_for_unit(gb, unit: str):
    if gb is None:
        return None
    low = _low(unit)
    if 'тера' in low or re.search(r'\bтб\b|\btb\b', low, re.I):
        return float(gb) / 1000.0
    return gb


def _usb_facts(plan: dict):
    row = _component(plan, 'Порты USB') or {}
    cons = dict(row.get('constraints') or {})
    keys = ('usb2_a','gen1_a','gen1_c','gen2_a','gen2_c','gen1_total','gen2_total','usb3x_total','type_c_total','front_usb2','front_gen1','front_a','front_c','rear_a')
    if all(k in cons for k in ('usb2_a','gen1_a','gen1_c','gen2_a','gen2_c')):
        return {k: int(cons.get(k, 0) or 0) for k in keys}
    # Backward-compatible parsing for projects saved by v1.0.x.
    name = _norm(row.get('name'))
    patterns = {
        'usb2_a': r'USB\s*2\.0\s*A:\s*(\d+)',
        'gen1_a': r'Gen1\s*A:\s*(\d+)',
        'gen1_c': r'Gen1\s*C:\s*(\d+)',
        'gen2_a': r'Gen2\+?\s*A:\s*(\d+)',
        'gen2_c': r'Gen2\+?\s*C:\s*(\d+)',
        'usb3x_total': r'USB\s*3\.x\s*всего:\s*(\d+)',
        'type_c_total': r'Type-C\s*всего:\s*(\d+)',
        'front_a': r'Type-A\s*спереди:\s*(\d+)',
        'front_c': r'Type-C\s*спереди:\s*(\d+)',
        'rear_a': r'Type-A\s*сзади:\s*(\d+)',
    }
    out = {}
    for k, pat in patterns.items():
        m = re.search(pat, name, re.I)
        if m:
            out[k] = int(m.group(1))
    out.setdefault('gen1_total', out.get('gen1_a',0)+out.get('gen1_c',0))
    out.setdefault('gen2_total', out.get('gen2_a',0)+out.get('gen2_c',0))
    out.setdefault('usb3x_total', out.get('gen1_total',0)+out.get('gen2_total',0))
    out.setdefault('type_c_total', out.get('gen1_c',0)+out.get('gen2_c',0))
    out.setdefault('front_usb2', 0)
    out.setdefault('front_gen1', 0)
    return out


def _case_facts(plan: dict):
    row = _component(plan, 'Корпус')
    if not row:
        return None
    cons = dict(row.get('constraints') or {})
    if cons.get('kind') == 'case':
        # v1.7+: the selected chassis carries its own factual specification so
        # custom cases do not inherit QX dimensions/ports/bays by name guessing.
        return cons
    name = _low(row.get('name'))
    if 'защищ' in name or 'тип 2' in name:
        return CASE_SPECS['protected']
    return CASE_SPECS['standard']


def _fmt_num(v: Any) -> str:
    if v is None:
        return ''
    try:
        f = float(v)
        return str(int(f)) if f.is_integer() else ('%.3f' % f).rstrip('0').rstrip('.')
    except Exception:
        return str(v)


def _aio_facts(plan: dict, req: dict):
    """Facts for the monoblock barebone actually used in the plan.

    Saved component constraints win. This is important when the specialist replaces
    the standard QX barebone with another platform: QX 23.8/27 facts must not leak
    into the application.
    """
    row = _component(plan, 'Моноблочная платформа') or {}
    cons = dict(row.get('constraints') or {})
    if cons.get('kind') == 'aio_platform':
        return cons
    # Backward compatibility for projects made before custom-platform support.
    if (req or {}).get('aio_platform_mode', 'qx') == 'custom':
        return {
            'custom': True,
            'name': (req or {}).get('custom_aio_name') or 'Другая моноблочная платформа',
            'diagonal_in': (req or {}).get('custom_aio_diagonal_in'),
            'resolution': (req or {}).get('custom_aio_resolution'),
            'refresh_hz': (req or {}).get('custom_aio_refresh_hz'),
            'matrix': (req or {}).get('custom_aio_matrix'),
            'built_in_cardreader': bool((req or {}).get('custom_aio_built_in_cardreader')),
            'integrated_psu': bool((req or {}).get('custom_aio_integrated_psu', True)),
            'psu_w': (req or {}).get('custom_aio_psu_w'),
        }
    return dict(AIO_SPECS.get((req or {}).get('product_type') or '', {}))


def _actual_for(name: str, req: dict, plan: dict, unit: str = '') -> Tuple[str, str]:
    low = _low(name)
    cpu = _selected_cpu(plan)
    platform = _platform(plan, req)
    board = _board_facts(plan, req)
    board_src = ('материнская плата из выбранной базы' if board.get('custom_profile') else 'материнская плата QX')
    ram = _ram_facts(plan)
    ssd = _ssd_facts(plan)
    hdd = _hdd_facts(plan)
    usb = _usb_facts(plan)
    case = _case_facts(plan)
    ptype = (req or {}).get('product_type') or ''
    aio = _aio_facts(plan, req) if str(ptype).startswith('aio') else {}

    # CPU
    if cpu:
        if 'тепловыделение процессора' in low:
            return _fmt_num(cpu.get('tdp_w')), 'выбранный процессор'
        if 'частота процессора максимальная' in low or 'максимальная частота процессора' in low:
            return _fmt_num(cpu.get('boost_ghz')), 'выбранный процессор'
        if 'частота процессора базовая' in low or 'базовая частота процессора' in low:
            return _fmt_num(cpu.get('base_ghz')), 'выбранный процессор'
        if 'количество потоков процессора' in low:
            return _fmt_num(cpu.get('threads')), 'выбранный процессор'
        if 'количество ядер процессора' in low:
            return _fmt_num(cpu.get('cores')), 'выбранный процессор'
        if 'кэш' in low and ('l3' in low or 'третьего уровня' in low):
            return _fmt_num(cpu.get('l3_mb')), 'выбранный процессор'
        if 'passmark' in low:
            return '', 'PassMark нужно проверять на дату подачи заявки'

    # RAM
    if 'тактовая частота оперативной памяти' in low or 'частота оперативной памяти' in low:
        return _fmt_num(ram.get('mhz')), 'выбранная оперативная память'
    if 'объем установленной оперативной памяти' in low or 'объем оперативной установленной памяти' in low:
        return _fmt_num(ram.get('gb')), 'выбранная оперативная память'
    if 'объем каждого установленного модуля оперативной памяти' in low:
        return _fmt_num(ram.get('module_gb')), 'выбранная оперативная память'
    if 'максимальный объем оперативной памяти' in low or 'максимальный объем поддерживаемой оперативной памяти' in low:
        return _fmt_num(board.get('ram_max_gb')), board_src
    if low == 'тип оперативной памяти':
        return ram.get('type') or '', 'выбранная оперативная память'

    # SSD/HDD
    if 'количество накопителей типа ssd' in low:
        return _fmt_num(ssd.get('count')), 'выбранная сборка'
    if 'общий объем накопителей ssd' in low:
        return _fmt_num(_drive_capacity_for_unit(ssd.get('total_gb'), unit)), 'выбранный SSD'
    if 'объем накопителя ssd' in low:
        return _fmt_num(_drive_capacity_for_unit(ssd.get('capacity_gb'), unit)), 'выбранный SSD'
    if 'интерфейс накопителя ssd' in low:
        return ssd.get('interface') or '', 'выбранный SSD'
    if 'количество накопителей типа hdd' in low:
        return _fmt_num(hdd.get('count')), 'выбранная сборка'
    if 'общий объем накопителей hdd' in low:
        return _fmt_num(_drive_capacity_for_unit(hdd.get('total_gb'), unit)), 'выбранный HDD'
    if 'объем накопителя hdd' in low:
        return _fmt_num(_drive_capacity_for_unit(hdd.get('capacity_gb'), unit)), 'выбранный HDD'

    # USB inventory
    if 'usb 3.2 gen 2' in low and 'type-c' in low:
        return _fmt_num(usb.get('gen2_c')), 'порты выбранной конфигурации'
    if 'usb 3.2 gen 2' in low and 'type-a' in low:
        return _fmt_num(usb.get('gen2_a')), 'порты выбранной конфигурации'
    if 'usb 3.2 gen 1' in low and 'type-c' in low:
        return _fmt_num(usb.get('gen1_c')), 'порты выбранной конфигурации'
    if 'usb 3.2 gen 1' in low and 'type-a' in low:
        return _fmt_num(usb.get('gen1_a')), 'порты выбранной конфигурации'
    if re.search(r'usb\s*3\.(?:x|х)', low):
        return _fmt_num(usb.get('usb3x_total')), 'порты выбранной конфигурации'
    # Location-specific requirements must be resolved before the generic Type-C total.
    if 'usb 2.0' in low and 'передн' in low:
        return _fmt_num(usb.get('front_usb2')), 'порты выбранной конфигурации'
    if 'usb type-a' in low and 'передн' in low:
        return _fmt_num(usb.get('front_a')), 'порты выбранной конфигурации'
    if 'usb type-c' in low and 'передн' in low:
        return _fmt_num(usb.get('front_c')), 'порты выбранной конфигурации'
    if 'usb type-a' in low and 'задн' in low:
        return _fmt_num(usb.get('rear_a')), 'порты выбранной конфигурации'
    if 'usb type-c' in low:
        return _fmt_num(usb.get('type_c_total')), 'порты выбранной конфигурации'
    if 'usb 2.0' in low:
        return _fmt_num(usb.get('usb2_a')), 'порты выбранной конфигурации'

    # Board / case
    if 'скорость передачи данных проводного сетевого контроллера' in low:
        return _fmt_num(board.get('ethernet_mbps')), board_src
    if ('8p8c' in low or 'rj-45' in low) and ('количество' in low or 'портов' in low) and ('сетев' in low or 'rj-45' in low):
        return _fmt_num(board.get('ethernet_ports')), board_src
    # Explicit port counts come from the selected board's video specification.
    # Do not assume that a custom-profile board has these outputs.
    video=board.get('video') or {}
    if ('количество портов hdmi' in low or 'количество разъемов hdmi' in low):
        return _fmt_num(video.get('HDMI')), board_src
    if ('количество' in low and ('displayport' in low or 'display port' in low) and ('портов' in low or 'порт' in low)):
        return _fmt_num(video.get('DisplayPort')), board_src
    if 'количество слотов m.2' in low and ('key m' in low or 'кеу m' in low):
        return _fmt_num(board.get('m2_key_m')), board_src
    if 'количество слотов m.2' in low and ('key e' in low or 'кеу e' in low):
        return _fmt_num(max(int(board.get('m2_key_e') or 0), int(board.get('m2_key_e_via_adapter') or 0))), board_src + (' + штатный переходник' if board.get('m2_key_e_via_adapter') else '')
    if 'количество слотов' in low and 'оперативной памяти' in low:
        return _fmt_num(board.get('ram_slots')), board_src
    if 'pci-e' in low and 'x1' in low:
        return _fmt_num(board.get('pcie_x1')), board_src
    if 'pci-e' in low and 'x16' in low:
        return _fmt_num(board.get('pcie_x16')), board_src
    case_src = 'другой корпус — вручную' if case and case.get('custom') else 'корпус QX'
    if low == 'ширина корпуса' and case:
        return _fmt_num(case.get('width_mm')), case_src
    if low == 'длина корпуса' and case:
        return _fmt_num(case.get('length_mm')), case_src
    if low == 'высота корпуса' and case:
        return _fmt_num(case.get('height_mm')), case_src
    if 'отсеков корпуса 5,25' in low and case:
        return _fmt_num(case.get('bay_525')), case_src
    if 'отсеков корпуса 2,5' in low and case:
        return _fmt_num(case.get('bay_25')), case_src
    if 'отсеков корпуса 3,5' in low and case:
        return _fmt_num(case.get('bay_35')), case_src

    # Display/AIO. Only use facts explicitly present in our QX specs.
    if 'разрешение экрана' in low and aio.get('resolution'):
        return str(aio['resolution']).replace('x', ' x '), ('другая моноблочная платформа — вручную' if aio.get('custom') else 'моноблочная платформа QX')
    if 'диагональ экрана' in low and aio.get('diagonal_in') is not None:
        return _fmt_num(aio.get('diagonal_in')), ('другая моноблочная платформа — вручную' if aio.get('custom') else 'моноблочная платформа QX')
    if 'частота обновления экрана' in low and aio.get('refresh_hz') is not None:
        return _fmt_num(aio.get('refresh_hz')), ('другая моноблочная платформа — вручную' if aio.get('custom') else 'моноблочная платформа QX')
    if 'технология изготовления матрицы' in low and aio.get('matrix'):
        return str(aio.get('matrix')), ('другая моноблочная платформа — вручную' if aio.get('custom') else 'моноблочная платформа QX')

    # PSU
    if 'мощность блока питания' in low:
        psu = _component(plan, 'Блок питания') or {}
        pc = psu.get('constraints') or {}
        if pc.get('w') is not None:
            return _fmt_num(pc.get('w')), 'встроенный БП выбранной платформы'
        m = re.search(r'(\d+)\s*Вт', _norm(psu.get('name')), re.I)
        if m:
            return m.group(1), 'выбранный блок питания'

    return '', 'нет подтверждённого фактического значения — заполни вручную'


def _extract_actual_number(value: str):
    # Accept ordinary decimals and spaced thousands such as «20 000».
    m = re.search(r'-?\d+(?:[ \u00a0]\d{3})*(?:[\.,]\d+)?', value or '')
    if not m:
        return None
    try:
        return float(m.group(0).replace(' ', '').replace('\u00a0','').replace(',', '.'))
    except Exception:
        return None


def requirement_status(requirement: str, actual: str) -> str:
    """Check factual value against the comparison operators written by the customer.

    Supports ≥, >, ≤, <, >=, <= and common Russian equivalents ("не менее",
    "не более", "более", "менее"). Strict operators remain strict. The function
    never blocks editing/export; it only returns a warning status for the UI.
    """
    req = _norm(requirement)
    actual = _norm(actual)
    if not actual:
        return 'НУЖНО ЗАПОЛНИТЬ'

    normalized = req
    # Normalize textual operators before shorter words to avoid "не менее" -> "не <".
    replacements = [
        (r'(?i)не\s+менее', '≥'),
        (r'(?i)не\s+более', '≤'),
        (r'(?i)больше\s+или\s+равно', '≥'),
        (r'(?i)меньше\s+или\s+равно', '≤'),
        (r'(?i)более', '>'),
        (r'(?i)менее', '<'),
    ]
    for pat, repl in replacements:
        normalized = re.sub(pat, repl, normalized)
    normalized = (normalized.replace('>=', '≥').replace('=>', '≥')
                  .replace('<=', '≤').replace('=<', '≤')
                  .replace('⩾', '≥').replace('⩽', '≤'))

    def cmp_num(a: float, op: str, b: float) -> bool:
        if op == '≥': return a >= b
        if op == '>': return a > b
        if op == '≤': return a <= b
        if op == '<': return a < b
        if op == '=': return abs(a - b) < 1e-9
        return True

    # Resolution/dimension pair, e.g. ≥2560 x 1440. Preserve strictness.
    rr = re.search(r'(≥|≤|>|<|=)?\s*(\d{2,5})\s*[xх×]\s*(\d{2,5})', normalized)
    ar = re.search(r'(\d{2,5})\s*[xх×]\s*(\d{2,5})', actual)
    if rr and ar:
        op = rr.group(1) or '='
        rw, rh = float(rr.group(2)), float(rr.group(3))
        aw, ah = float(ar.group(1)), float(ar.group(2))
        ok = cmp_num(aw, op, rw) and cmp_num(ah, op, rh)
        return 'OK' if ok else 'НЕ СООТВЕТСТВУЕТ'

    a = _extract_actual_number(actual)
    if a is not None:
        comparisons = [(m.group(1), float(m.group(2).replace(' ', '').replace('\u00a0','').replace(',', '.')))
                       for m in re.finditer(r'(≥|≤|>|<|=)\s*(-?\d+(?:[ \u00a0]\d{3})*(?:[\.,]\d+)?)', normalized)]
        if comparisons:
            ok = all(cmp_num(a, op, b) for op, b in comparisons)
            return 'OK' if ok else 'НЕ СООТВЕТСТВУЕТ'

    # Exact textual/categorical rows are normally locked by the customer. If an editable
    # row contains no machine-readable comparison, keep it user-editable without guessing.
    return 'ЗАПОЛНЕНО'


def build_characteristic_rows(characteristics: Iterable[dict], reqs: List[dict], plans: List[dict], overrides: Dict[str, str] | None = None) -> List[dict]:
    overrides = overrides or {}
    out = []
    for ch in characteristics or []:
        idx = int(ch.get('plan_index') or 0)
        req = reqs[idx] if idx < len(reqs) else (reqs[0] if reqs else {})
        plan = plans[idx] if idx < len(plans) else (plans[0] if plans else {})
        if ch.get('editable'):
            auto, source = _actual_for(ch.get('name',''), req, plan, ch.get('unit',''))
            auto_status = requirement_status(ch.get('original_value',''), auto) if auto else 'НУЖНО ЗАПОЛНИТЬ'
            if ch['id'] in overrides:
                final = str(overrides.get(ch['id']) or '')
                status = requirement_status(ch.get('original_value',''), final)
            elif auto and auto_status == 'НЕ СООТВЕТСТВУЕТ':
                # Never silently insert a known non-compliant auto value into an application.
                # The specialist may still type any value manually; then we only warn.
                final = ''
                status = 'АВТО НЕ СООТВЕТСТВУЕТ'
            else:
                final = str(auto or '')
                status = auto_status
        else:
            auto, source = ch.get('original_value',''), 'зафиксировано заказчиком'
            auto_status = 'НЕ ИЗМЕНЯЕТСЯ'
            final = ch.get('original_value','')
            status = 'НЕ ИЗМЕНЯЕТСЯ'
        item = dict(ch)
        item.update({'auto_value': auto, 'auto_status': auto_status, 'value': final, 'source': source, 'status': status})
        out.append(item)
    return out


def _replace_cell_text(cell, text: str):
    """Replace visible cell text while preserving the cell's existing paragraph/run formatting."""
    text = str(text)
    if not cell.paragraphs:
        cell.add_paragraph(text)
        return
    first = cell.paragraphs[0]
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run.text = ''
    else:
        first.add_run(text)
    for p in cell.paragraphs[1:]:
        for run in p.runs:
            run.text = ''


def fill_docx_characteristics(data: bytes, characteristics: Iterable[dict], values: Dict[str, str]) -> bytes:
    doc = Document(BytesIO(data))
    tables = list(doc.tables)
    for ch in characteristics or []:
        if not ch.get('editable'):
            continue
        value = _norm(values.get(ch.get('id'), ''))
        if not value:
            continue
        ti = int(ch['table_index']); ri = int(ch['row_index']); ci = int(ch['value_col'])
        if 0 <= ti < len(tables) and 0 <= ri < len(tables[ti].rows) and 0 <= ci < len(tables[ti].rows[ri].cells):
            _replace_cell_text(tables[ti].rows[ri].cells[ci], value)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
