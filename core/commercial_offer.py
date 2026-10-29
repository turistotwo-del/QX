from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
import json
from pathlib import Path

from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = BASE_DIR / 'templates' / 'kp_template.docx'
LOGO_FILE = BASE_DIR / 'templates' / 'logo_qx.jpeg'
SEQUENCE_FILE = BASE_DIR / 'data' / 'kp_sequence.json'


def money_int(value: float) -> str:
    return f'{int(round(float(value or 0))):,}'.replace(',', ' ')


def outgoing_number(d: date | datetime | None = None, sequence: int = 1) -> str:
    if d is None:
        d = date.today()
    if isinstance(d, datetime):
        d = d.date()
    return f'БрП-{d:%d%m%y}-{int(sequence)}'


def _load_sequence_state() -> dict:
    try:
        obj = json.loads(SEQUENCE_FILE.read_text(encoding='utf-8'))
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}


def next_sequence(d: date | None = None) -> int:
    d = d or date.today()
    state = _load_sequence_state()
    key = d.isoformat()
    return int(state.get(key, 0) or 0) + 1


def reserve_sequence(d: date | None, seq: int) -> None:
    d = d or date.today()
    SEQUENCE_FILE.parent.mkdir(parents=True, exist_ok=True)
    state = _load_sequence_state()
    key = d.isoformat()
    state[key] = max(int(state.get(key, 0) or 0), int(seq))
    # Keep the file small.
    items = sorted(state.items(), reverse=True)[:120]
    SEQUENCE_FILE.write_text(json.dumps(dict(items), ensure_ascii=False, indent=2), encoding='utf-8')


def calculate_offer_price(unit_cost: float, percent: float, mode: str = 'markup') -> float:
    unit_cost = float(unit_cost or 0)
    pct = max(0.0, float(percent or 0)) / 100.0
    if mode == 'margin':
        if pct >= 0.999999:
            raise ValueError('Маржа должна быть меньше 100%')
        return unit_cost / (1.0 - pct)
    return unit_cost * (1.0 + pct)


def build_offer_items(plans: list[dict], percent: float, mode: str = 'markup') -> list[dict]:
    items = []
    for idx, plan in enumerate(plans or [], start=1):
        req = plan.get('req') or {}
        qty = int(req.get('quantity') or 0)
        if qty <= 0:
            continue
        rows = plan.get('rows') or []
        unit_cost = sum(float(r.get('qty') or 0) * float(r.get('price') or 0) for r in rows)
        # КП template shows whole rubles. Round the unit selling price first so that
        # displayed unit price × quantity always equals the displayed line total.
        price = float(round(calculate_offer_price(unit_cost, percent, mode)))
        reg = plan.get('registry') or {}
        name = str(reg.get('product_name') or req.get('item_name') or req.get('title') or f'Позиция {idx}')
        registry_number = str(reg.get('registry_number') or '')
        items.append({
            'name': name,
            'registry_number': registry_number,
            'quantity': qty,
            'unit_cost': unit_cost,
            'unit_price': price,
            'total': price * qty,
        })
    return items



def _remove_table_borders(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'nil')


def _rebuild_company_header(doc):
    """Create a stable two-column header aligned with the document's left margin.
    No nested table (Word used to relocate it when the document opened).
    """
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    if not doc.tables or not LOGO_FILE.exists():
        return
    outer=doc.tables[0]
    outer.autofit=False
    outer.alignment=WD_TABLE_ALIGNMENT.LEFT
    from docx.oxml.ns import qn
    indent = outer._tbl.tblPr.find(qn('w:tblInd'))
    if indent is not None: indent.set(qn('w:w'), '0')
    if len(outer.columns)==1:
        outer.add_column(Inches(4.55))
    outer.columns[0].width=Inches(1.55)
    outer.columns[1].width=Inches(4.70)
    _remove_table_borders(outer)
    left,right=outer.rows[0].cells[:2]
    left.width=Inches(1.55)
    right.width=Inches(4.70)
    for cell in (left,right):
        for t in list(cell._tc):
            if not t.tag.endswith('}tcPr'): cell._tc.remove(t)
        cell.add_paragraph('')
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    lp=left.paragraphs[0]
    lp.alignment=WD_ALIGN_PARAGRAPH.LEFT
    lp.add_run().add_picture(str(LOGO_FILE),width=Inches(1.32))
    rp=right.paragraphs[0]
    rp.paragraph_format.space_after=Pt(0)
    rp.paragraph_format.space_before=Pt(0)
    lines=[
       ('Общество с ограниченной ответственностью',True),
       ('«БРОКТЕХНОЛОДЖИ»',True),
       ('ИНН/КПП 5257201533/525701001; ОГРН',False),
       ('1215200005900',False),
       ('603002, Нижегородская область, г. Нижний Новгород,',False),
       ('ул. Обухова, д. 11, офис 305/1, тел.: 8 915 957 69 98',False)]
    for i,(txt,emphasis) in enumerate(lines):
        run=rp.add_run(txt)
        run.font.name='Times New Roman'
        run.font.size=Pt(10.5)
        run.font.bold=emphasis
        run.font.italic=emphasis
        run.font.color.rgb=RGBColor(115,115,115)
        if i<len(lines)-1:run.add_break()


def _set_paragraph_text(paragraph, text: str) -> None:
    # Keep paragraph-level formatting/alignment and reuse the formatting of the first run.
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)


def _set_cell_text(cell, text: str) -> None:
    # Preserve cell/table geometry; replace text in first paragraph only.
    if not cell.paragraphs:
        cell.text = str(text)
        return
    _set_paragraph_text(cell.paragraphs[0], str(text))
    for p in cell.paragraphs[1:]:
        _set_paragraph_text(p, '')


def _insert_product_rows(table, count: int):
    if len(table.rows) < 3:
        raise ValueError('Шаблон КП повреждён: нет строки товара/итога')
    product_tr = table.rows[1]._tr
    total_tr = table.rows[-1]._tr
    # Remove any pre-existing extra item rows, leaving header + one item + total.
    while len(table.rows) > 3:
        table._tbl.remove(table.rows[-2]._tr)
    rows = [table.rows[1]]
    for _ in range(max(0, count - 1)):
        new_tr = deepcopy(product_tr)
        total_tr.addprevious(new_tr)
        rows.append(table.rows[-2])
    return rows, table.rows[-1]


def build_commercial_offer(
    items: list[dict],
    *,
    sequence: int = 1,
    offer_date: date | datetime | None = None,
    warranty_months: int = 12,
    validity_days: int = 30,
    outgoing: str | None = None,
    template_path: str | Path | None = None,
) -> bytes:
    if not items:
        raise ValueError('Для КП нет товарных позиций с количеством больше нуля')
    if offer_date is None:
        offer_date = date.today()
    if isinstance(offer_date, datetime):
        offer_date = offer_date.date()
    outgoing = outgoing or outgoing_number(offer_date, sequence)
    template = Path(template_path) if template_path else DEFAULT_TEMPLATE
    if not template.exists():
        raise FileNotFoundError(f'Не найден шаблон КП: {template}')

    doc = Document(str(template))
    _rebuild_company_header(doc)
    # Converted sample has these paragraphs. Find by prefix rather than fixed index where possible.
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith('Исходящий номер:'):
            _set_paragraph_text(p, f'Исходящий номер: {outgoing}')
        elif txt.startswith('Дата письма:'):
            _set_paragraph_text(p, f'Дата письма: {offer_date:%d.%m.%Y} г.')
        elif txt.startswith('Гарантия на изделие'):
            _set_paragraph_text(p, f'Гарантия на изделие – {int(warranty_months)} месяцев.')
        elif txt.startswith('Срок действия предложения'):
            _set_paragraph_text(p, f'Срок действия предложения – {int(validity_days)} дней.')

    if len(doc.tables) < 2:
        raise ValueError('Шаблон КП повреждён: не найдена таблица товаров')
    table = doc.tables[1]
    product_rows, total_row = _insert_product_rows(table, len(items))
    grand_total = 0.0
    for i, (row, item) in enumerate(zip(product_rows, items), start=1):
        name = str(item.get('name') or '')
        reg = str(item.get('registry_number') or '').strip()
        if reg:
            name += f'\nРеестровый номер: {reg}'
        qty = int(item.get('quantity') or 0)
        unit = float(item.get('unit_price') or 0)
        total = float(item.get('total') if item.get('total') is not None else unit * qty)
        grand_total += total
        vals = [str(i), name, f'{qty}\nшт.', money_int(unit), money_int(total)]
        for c, v in zip(row.cells, vals):
            _set_cell_text(c, v)
    _set_cell_text(total_row.cells[-1], money_int(grand_total))

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
