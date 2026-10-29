"""Offline acceptance for multi-file EIS intake and all-part Word evidence."""
from __future__ import annotations

import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

os.environ.setdefault('QX_DATA_DIR', tempfile.mkdtemp(prefix='qx_eis_batch_boot_'))

from docx import Document
from fastapi.testclient import TestClient
from modern import tender_store as store
from modern.eis_document_download import import_official_documents
from modern.tender_analysis import document_segments
from modern.server import app


def make_word(text, header='', footer=''):
    doc = Document()
    doc.add_paragraph(text)
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    if footer:
        doc.sections[0].footer.paragraphs[0].text = footer
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


class _FakeResponse:
    def __init__(self, url, data):
        self.url = url
        self.stream = BytesIO(data)
        self.headers = {'Content-Type':'application/octet-stream'}
    def __enter__(self): return self
    def __exit__(self, *_): return False
    def geturl(self): return self.url
    def read(self, n=-1): return self.stream.read(n)


class _FakeOpener:
    def __init__(self, docs): self.docs=docs
    def open(self, request, timeout):
        if request.full_url not in self.docs:
            raise ValueError('Недоступен документ')
        return _FakeResponse(request.full_url, self.docs[request.full_url])


class TestMultiEisIntake(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=TemporaryDirectory(prefix='qx_batch_eis_test_')
        cls.old=(store.APP_DATA,store.DB_FILE,store.DOCUMENTS,store._INITIALIZED)
        store.APP_DATA=Path(cls.tmp.name)
        store.DB_FILE=store.APP_DATA/'tenders.sqlite3'
        store.DOCUMENTS=store.APP_DATA/'tender_documents'
        store._INITIALIZED=False
        store.initialize()
        cls.client=TestClient(app)

    @classmethod
    def tearDownClass(cls):
        store.APP_DATA,store.DB_FILE,store.DOCUMENTS,store._INITIALIZED=cls.old
        cls.tmp.cleanup()

    def _tid(self): return store.create({'title':'Проверка пачки документов'})['id']

    def test_two_official_docs_and_sources(self):
        tid=self._tid()
        a='https://zakupki.gov.ru/docs/contract.docx'
        b='https://zakupki.gov.ru/docs/technical.docx'
        reply=import_official_documents(tid,[a,b],opener=_FakeOpener({
            a:make_word('Поставка осуществляется в течение 20 календарных дней с даты заключения контракта'),
            b:make_word('Требование к реестровому BIOS обязательно'),
        }))
        self.assertEqual(reply['success_count'],2)
        self.assertEqual(reply['failed_count'],0)
        self.assertEqual(len(store.detail(tid)['documents']),2)
        self.assertNotIn(a, str(store.detail(tid)))

    def test_invalid_url_rejects_entire_batch_before_network(self):
        tid=self._tid()
        before=len(store.detail(tid)['documents'])
        with self.assertRaises(ValueError):
            import_official_documents(tid,[
                'https://zakupki.gov.ru/docs/valid.docx',
                'https://evil.example/private.docx',
            ],opener=_FakeOpener({}))
        self.assertEqual(len(store.detail(tid)['documents']),before)

    def test_partial_success_is_explicit(self):
        tid=self._tid()
        good='https://zakupki.gov.ru/docs/first.docx'
        missing='https://zakupki.gov.ru/docs/missing.docx'
        reply=import_official_documents(tid,[good,missing],opener=_FakeOpener({
            good:make_word('Техническое задание системы')
        }))
        self.assertEqual(reply['success_count'],1)
        self.assertEqual(reply['failed_count'],1)
        self.assertEqual(reply['failures'][0]['line'],2)
        self.assertEqual(len(store.detail(tid)['documents']),1)

    def test_repeated_or_too_many_links(self):
        tid=self._tid()
        good='https://zakupki.gov.ru/a.docx'
        for entries in ([good,good],[],[good+f'?i={i}' for i in range(11)]):
            with self.assertRaises(ValueError):
                import_official_documents(tid,entries,opener=_FakeOpener({}))
        self.assertEqual(len(store.detail(tid)['documents']),0)

    def test_no_token_leak_in_error(self):
        tid=self._tid()
        sensitive='https://zakupki.gov.ru/secret.docx?token=VERY_PRIVATE'
        with patch('modern.eis_document_download.import_official_document', side_effect=ValueError(sensitive)):
            reply=import_official_documents(tid,[sensitive])
        self.assertEqual(reply['failed_count'],1)
        self.assertNotIn('VERY_PRIVATE',repr(reply))

    def test_http_endpoint(self):
        tid=self._tid()
        urls=['https://zakupki.gov.ru/files/a.docx','https://zakupki.gov.ru/files/b.docx']
        with patch('modern.eis_document_download.import_official_document') as func:
            func.return_value={'imported':['a.docx'],'warnings':[],'duplicates':[]}
            response=self.client.post(f'/api/tenders/{tid}/eis/documents-from-urls', json={'urls':urls})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['success_count'],2)

    def test_http_rejects_url_attack(self):
        tid=self._tid()
        response=self.client.post(f'/api/tenders/{tid}/eis/documents-from-urls',json={
            'urls':['https://zakupki.gov.ru/a.docx','http://127.0.0.1/private.docx']})
        self.assertEqual(response.status_code,422)


class TestWordExtraParts(unittest.TestCase):
    def test_header_clause_available_with_precise_location(self):
        data=make_word('Основной текст',header='Реестровый BIOS обязателен для оборудования')
        chunks=document_segments('Контракт.docx',data)
        self.assertTrue(any('верхний колонтитул' in x['location'] and 'Реестровый BIOS' in x['text'] for x in chunks))

    def test_footer_clause_available_with_precise_location(self):
        data=make_word('Основной текст',footer='Гарантийный срок 36 месяцев')
        chunks=document_segments('Контракт.docx',data)
        self.assertTrue(any('нижний колонтитул' in x['location'] and 'Гарантийный срок' in x['text'] for x in chunks))


if __name__=='__main__': unittest.main()
