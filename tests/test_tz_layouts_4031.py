"""Real document-layout fixtures generated offline, no LLM or network required."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from unittest import TestCase
from unittest.mock import patch
import json
from docx import Document
from openpyxl import Workbook
from core.tz_parser import parse_tz, extract_text
from modern import local_ai, tz_ai


def docx_reversed_headers():
    d = Document();d.add_paragraph('Техническое задание по закупке компьютерной техники')
    tbl = d.add_table(rows=1,cols=4)
    for c,value in zip(tbl.rows[0].cells, ('Требуемое значение','Инструкция по заполнению', 'Наименование характеристики','Единица измерения')):
        c.text = value
    for row in [('≥ 16','конкретное значение','Объем установленной оперативной памяти','Гигабайт'),
                ('≥ 480','','Общий объем накопителей SSD','Гигабайт')]:
        for c,value in zip(tbl.add_row().cells,row):c.text=value
    b=BytesIO();d.save(b);return b.getvalue()


class Layout4031Tests(TestCase):
    def test_docx_kt_ru_reordered_columns(self):
        reqs,text=parse_tz('Переставленные_столбцы.docx',docx_reversed_headers())
        self.assertEqual(len(reqs),1)
        self.assertEqual(reqs[0].ram_gb_min,16)
        self.assertEqual(reqs[0].ssd_gb_min,480)
        self.assertIn('Наименование характеристики',text)

    def test_excel_per_sheet_characteristics(self):
        book=Workbook();first=book.active;first.title='Моноблок 23.8'
        first.append(('Значение характеристики','Наименование характеристики','Единица измерения'))
        first.append(('≥ 16','Объем установленной оперативной памяти','Гигабайт'))
        second=book.create_sheet('Монитор 27')
        second.append(('Характеристика','Значение'))
        second.append(('Диагональ экрана','≥ 27'))
        stream=BytesIO();book.save(stream)
        reqs,_=parse_tz('две_позиции.xlsx',stream.getvalue())
        self.assertEqual(len(reqs),2)
        self.assertEqual([r.product_type for r in reqs],['aio24','monitor'])
        self.assertEqual(reqs[0].ram_gb_min,16)
        self.assertEqual(reqs[1].monitor_inches,27)

    def test_excel_procurement_quantity_is_not_characteristic(self):
        wb=Workbook();s=wb.active;s.title='Спецификация'
        s.append(['Наименование товара','Описание технических характеристик','Количество'])
        s.append(['Моноблок 23,8', 'Оперативная память: не менее 16 ГБ; SSD: 480 ГБ', 19])
        s.append(['Системный блок', 'Объем ОЗУ: не менее 8 ГБ', 35])
        b=BytesIO();wb.save(b);reqs,_=parse_tz('товары.xlsx',b.getvalue())
        self.assertEqual(len(reqs),2)
        self.assertEqual([r.quantity for r in reqs],[19,35])
        self.assertEqual(reqs[0].product_type,'aio24')
        self.assertEqual(reqs[1].ram_gb_min,8)

    def test_txt_colon_alias_and_positions(self):
        data=('Позиция №1 Системный блок\nОбъем ОЗУ: от 16 ГБ\n'
              'Число ядер — не менее 6\nПозиция №2 Монитор\n'
              'Диагональ экрана: не менее 27 дюймов').encode()
        reqs,_=parse_tz('ТЗ_произвольное.txt',data)
        self.assertEqual(len(reqs),2)
        self.assertEqual(reqs[0].ram_gb_min,16)
        self.assertEqual(reqs[0].cpu_cores_min,6)
        self.assertEqual(reqs[1].product_type,'monitor')
        self.assertEqual(reqs[1].monitor_inches,27)

    def test_odt_without_external_program(self):
        payload='''<?xml version="1.0" encoding="utf-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text><text:p>Монитор 27</text:p><text:p>Диагональ экрана: 27 дюймов</text:p></office:text></office:body></office:document-content>'''.encode('utf-8')
        b=BytesIO()
        with ZipFile(b,'w',ZIP_DEFLATED) as z:z.writestr('content.xml',payload)
        reqs,txt=parse_tz('монитор.odt',b.getvalue())
        self.assertIn('Диагональ экрана',txt)
        self.assertEqual(reqs[0].monitor_inches,27)

    def test_long_document_retrieves_very_late_evidence(self):
        data='Общие условия согласования и поставки.\n'*2000+'\nОперативная память не менее 32 ГБ'
        prompt,truncated=tz_ai._lite_prompt(data, {'product_type':'system','title':'Системный блок'})
        self.assertTrue(truncated)
        self.assertIn('Оперативная память не менее 32 ГБ',prompt)
        self.assertLess(len(prompt),5500)

    def test_llm_can_fill_only_safe_empty_numeric_values(self):
        src='Оперативная память не менее 16 ГБ.\nВысота корпуса не более 400 мм'
        output=tz_ai.verify_ai_proposals({'proposals':[
          {'field':'ram_gb_min','value':16,'quote':'Оперативная память не менее 16 ГБ'},
          {'field':'case_height_max_mm','value':400,'quote':'Высота корпуса не более 400 мм'},
        ]},src,{'ram_gb_min':None,'case_height_max_mm':None})
        by_key={p['field']:p for p in output['proposals']}
        self.assertTrue(by_key['ram_gb_min']['auto_safe'])
        self.assertFalse(by_key['case_height_max_mm']['auto_safe'])
        edited=tz_ai.verify_ai_proposals({'proposals':[
          {'field':'ram_gb_min','value':16,'quote':'Оперативная память не менее 16 ГБ'}
        ]},src,{'ram_gb_min':8})
        self.assertFalse(edited['proposals'][0]['auto_safe'])

    def test_reject_cross_position_quote_even_if_document_contains_it(self):
        positions='Позиция 1 SSD 256 ГБ\nПозиция 2 SSD 512 ГБ'
        req={'title':'Позиция 1', 'raw_text':'Позиция 1 SSD 256 ГБ','ssd_gb_min':None}
        raw=json.dumps({'proposals':[{'field':'ssd_gb_min','value':512,'quote':'Позиция 2 SSD 512 ГБ'}]})
        with patch.object(local_ai,'available_models',return_value={'models':[local_ai.LITE_MODEL]}), \
             patch.object(local_ai,'_lite_answer',return_value=raw):
            result=tz_ai.analyze_tz(positions,req,local_ai.LITE_MODEL)
        self.assertFalse(result['proposals'])

    def test_ui_autofill_and_multi_position_controls_exist(self):
        from pathlib import Path
        js=(Path(__file__).resolve().parents[1]/'modern/web/app.js').read_text()
        self.assertIn('tz-ai-all',js)
        self.assertIn('tz-ai-fill',js)
        self.assertIn('auto_safe === true',js)
        self.assertIn('runTzAiBatch',js)
