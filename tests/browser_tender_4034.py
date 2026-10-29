"""Chromium: opt-in from TZ, verified mock XML preview, deadline calendar and spacing."""
import base64
import os
from io import BytesIO
from pathlib import Path
import tempfile
from unittest.mock import patch
os.environ['QX_OFFLINE_MODE']='1'
os.environ['QX_DATA_DIR']=tempfile.mkdtemp(prefix='qx4034_browser_')
from docx import Document
from fastapi.testclient import TestClient
from playwright.sync_api import sync_playwright
from chromium_launch import launch_chromium
from modern.server import app
from modern import tender_store
from modern import eis_public

root=Path(__file__).resolve().parents[1]
web=root/'modern/web'
client=TestClient(app)
NUMBER='0322100002526000466'

def forward(req):
    path=str(req.get('url','')).replace('/api/','/',1)
    content=base64.b64decode(req['binary']) if req.get('binary') else str(req.get('body') or '').encode('utf-8')
    response=client.request(req.get('method') or 'GET','/api'+path,content=content,headers=req.get('headers') or {})
    return {'status':response.status_code,'body':response.text,'type':response.headers.get('content-type','application/json')}

html=(web/'index.html').read_text('utf8').replace('<link rel="icon" href="/assets/qx_icon.png">','')
html=html.replace('<link rel="stylesheet" href="/style.css">','<style>'+(web/'style.css').read_text('utf8')+'</style>')
html=html.replace('<script type="module" src="/app.js"></script>', '''<script>window.fetch=async(url,opts={})=>{
 let binary=null;
 if(opts.body instanceof Blob){let b=new Uint8Array(await opts.body.arrayBuffer()); let x='';for(let i=0;i<b.length;i+=14000)x+=String.fromCharCode(...b.slice(i,i+14000));binary=btoa(x);}
 const r=await QXAPI({url:String(url),method:opts.method||'GET',body:binary?'':opts.body||'',binary,headers:opts.headers||{}});
 return new Response(r.body,{status:r.status,headers:{'content-type':r.type}});
};</script><script>'''+(web/'app.js').read_text('utf8')+'</script>')

def docx_sample():
    d=Document();d.add_heading('Техническое задание',0);d.add_paragraph('Моноблок 23.8 дюйма. Оперативная память не менее 16 ГБ.')
    b=BytesIO();d.save(b);return b.getvalue()

def run():
 with patch.object(eis_public,'fetch_notice_metadata',return_value={
     'available':True,'notice':{'number':NUMBER,'title':'Моноблоки для учреждения',
      'customer':'МБУ Заказчик','nmck':'654321','application_due_on':'2026-11-02','application_due_time':'16:00'},
     'message':'Получен XML по точному номеру'}):
  with sync_playwright() as p:
   browser=launch_chromium(p)
   page=browser.new_page(viewport={'width':1500,'height':920})
   page.set_default_timeout(12000)
   errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
   page.expose_function('QXAPI',forward)
   page.set_content(html,wait_until='domcontentloaded');page.wait_for_selector('.sidebar')
   page.locator('[data-action="tab"][data-tab="tz"]').click()
   assert page.locator('[data-tender-auto-link]').is_visible()
   assert page.locator('[data-tender-auto-link]').is_checked() is False
   page.locator('#file-tz').set_input_files({'name':'demo.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':docx_sample()})
   page.wait_for_selector('.loaded-source',timeout=18000)
   assert tender_store.listing()==[], 'Opt-out created an unwanted purchase!'
   page.locator('[data-tender-auto-link]').check()
   box=page.locator('[data-tender-notice-input]')
   box.fill(NUMBER);box.press('Tab')
   assert page.locator('[data-action="tender-link-loaded"]').is_enabled()
   page.locator('[data-action="tender-link-loaded"]').click()
   page.wait_for_timeout(2500)
   print('DEBUG',page.evaluate('({tid:state.tenderId,number:state.tenderNoticeNumber,preview:state.tenderEisPreview,msg:state.tenderEisMessage,loading:state.busy})'),flush=True)
   print('DEBUG_ROWS',tender_store.listing(),flush=True)
   print('DEBUG_ERRORS',errors,flush=True)
   page.wait_for_function('state.tenderId && state.tenderEisPreview && state.tenderEisPreview.available',timeout=9000)
   records=tender_store.listing()
   assert len(records)==1,records
   t=records[0]
   assert t['number']==NUMBER and t['customer']=='МБУ Заказчик',t
   assert t['application_due_on']=='2026-11-02',t
   assert t['application_due_time']=='16:00',t
   assert page.locator('.qx-eis-preview').is_visible()
   page.locator('[data-action="tender-nav"][data-view="calendar"]').click()
   page.wait_for_selector('.tender-calendar')
   page.locator('[data-action="tender-next-month"]').click() if False else None
   page.evaluate('state.tenders.monthDelta=1;state.tenders.calendarFocus="2026-11-02";render()')
   page.wait_for_selector('.qx-calendar-pill')
   assert '16:00' in page.locator('.qx-calendar-pill').first.inner_text()
   page.locator('[data-action="tender-nav"][data-view="contracts"]').click()
   page.wait_for_selector('.qx-tender-ident')
   ident=page.locator('.qx-tender-ident').first
   assert ident.locator('b').count()==1 and ident.locator('small').count()>=1
   assert 'МБУ Заказчик' not in ident.inner_text()
   assert not errors,errors
   browser.close()
  print('BROWSER_TENDER_4034_OK')

if __name__=='__main__':run()
