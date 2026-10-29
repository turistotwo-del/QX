"""No network/model downloads: check real response schemas and safety gates."""
from __future__ import annotations
import asyncio
import base64
from io import BytesIO
import json
import tempfile
import unittest
from unittest.mock import patch
import httpx
from docx import Document
from modern.server import app, project_dump
from modern.tz_ai import verify_ai_proposals, analyze_tz
from core.project_io import load_project
from modern import tz_ai, local_ai


def docx_blob(line):
    doc = Document()
    doc.add_heading('Техническое задание', level=1)
    doc.add_paragraph(line)
    buf=BytesIO();doc.save(buf);return buf.getvalue()


class TZMultiDocTests(unittest.TestCase):
    def test_ai_rejects_invented_fields_values_and_citations(self):
        source='Оперативная память не менее 16 ГБ. USB Type-C не менее 2 портов.'
        result=verify_ai_proposals({'proposals':[
            {'field':'ram_gb_min','value':16,'quote':'Оперативная память не менее 16 ГБ'},
            {'field':'need_gpu','value':True,'quote':'дискретная видеокарта не менее 8 ГБ'},
            {'field':'psu_w_min','value':900,'quote':'дискретная видеокарта не менее 8 ГБ'},
            {'field':'ram_gb_min','value':99999,'quote':'Оперативная память не менее 16 ГБ'},
            {'field':'equipment_profile_name','value':'Заменить прайс','quote':'Оперативная память не менее 16 ГБ'},
            {'field':'usb_c_total_min','value':2,'quote':'USB Type-C не менее 2 портов'},
            {'field':'need_gpu','value':'true','quote':'USB Type-C не менее 2 портов'},
            {'field':'ram_type_min_generation','value':True,'quote':'Оперативная память не менее 16 ГБ'},
        ]},source,{'ram_gb_min':8})
        self.assertEqual([(r['field'],r['value']) for r in result['proposals']],
                         [('ram_gb_min',16),('usb_c_total_min',2)])
        self.assertGreaterEqual(result['rejected'],4)
        mismatch=verify_ai_proposals({'proposals':[{'field':'ram_gb_min','value':32,
            'quote':'Оперативная память не менее 16 ГБ'}]},source,{'ram_gb_min':8})
        self.assertEqual(mismatch['proposals'],[])

    def test_real_model_adapter_mocked_only_on_local_endpoint(self):
        class Reply:
            def __enter__(self):return self
            def __exit__(self,*args):pass
            def read(self,*args):return json.dumps({'response':json.dumps({'proposals':[
                {'field':'ram_gb_min','value':32,'quote':'Оперативная память не менее 32 ГБ'},
            ]},ensure_ascii=False)}).encode()
        with patch.object(local_ai,'available_models',return_value={'models':['qwen2.5:3b']}), \
             patch.object(tz_ai,'urlopen',side_effect=lambda request,timeout=None: Reply()) as mock_post:
            result=analyze_tz('Оперативная память не менее 32 ГБ',
                              {'title':'Системный блок','ram_gb_min':16},'qwen2.5:3b')
        self.assertEqual(result['proposals'][0]['value'],32)
        self.assertTrue(mock_post.call_args.args[0].full_url.startswith('http://127.0.0.1:11434/'))
        self.assertFalse(result.get('applied',False))

    def test_ai_requires_model(self):
        with self.assertRaises(ValueError):
            analyze_tz('Достаточно текста',{},'offline-search')

    def test_parse_two_documents_and_project_roundtrip(self):
        async def run():
            a=docx_blob('Системный блок, оперативная память не менее 8 ГБ.')
            b=docx_blob('Моноблок, оперативная память не менее 16 ГБ.')
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://testserver') as client:
                documents=[];reqs=[]
                for i,(name,data) in enumerate([('system.docx',a),('aio.docx',b)]):
                    response=await client.post('/api/parse',content=data,headers={'X-File-Name':name})
                    self.assertEqual(response.status_code,200,response.text)
                    result=response.json()
                    self.assertTrue(result['reqs'])
                    documents.append({'id':str(i),'name':name,'text':result['source_text'],
                                      'docx':result['source_docx_base64'],
                                      'characteristics':result['characteristics']})
                    reqs += [dict(r,_qx_doc_id=str(i)) for r in result['reqs']]
                project=project_dump({'reqs':reqs,'plans':[],'tz_documents':documents,
                                      'name':'ТЗ №1 + ТЗ №2'})
                loaded=load_project(project)
                self.assertEqual(len(loaded['tz_documents']),2)
                self.assertEqual({r['_qx_doc_id'] for r in loaded['reqs']},{'0','1'})
                self.assertTrue(base64.b64decode(loaded['tz_documents'][1]['docx']).startswith(b'PK'))
                resp=await client.post('/api/parse/ai',json={'source_text':'abc','req':{},'model':'offline-search'})
                self.assertEqual(resp.status_code,422)
        asyncio.run(run())

    def test_health_not_blocked_by_ai(self):
        async def run():
            def slow(*args):
                import time;time.sleep(.35);return {'proposals':[]}
            with patch.object(tz_ai,'analyze_tz',side_effect=slow):
                async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://testserver') as client:
                    task=asyncio.create_task(client.post('/api/parse/ai',json={'model':'m','source_text':'x','req':{}}))
                    await asyncio.sleep(.05)
                    resp=await asyncio.wait_for(client.get('/api/health'),timeout=.2)
                    self.assertEqual(resp.status_code,200)
                    self.assertEqual((await task).status_code,200)
        asyncio.run(run())

if __name__=='__main__':unittest.main()
