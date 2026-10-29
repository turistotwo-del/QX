"""Importing untrusted user-provided EIS documentation ZIP bundles."""
import io
import os
from pathlib import Path
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

os.environ.setdefault('QX_DATA_DIR', tempfile.mkdtemp(prefix='qx_package_boot_'))
from fastapi.testclient import TestClient
from docx import Document
from modern import tender_store as store
from modern.server import app


def archive(files):
    buf=io.BytesIO()
    with ZipFile(buf,'w',ZIP_DEFLATED) as z:
        for name, contents in files:
            z.writestr(name,contents)
    return buf.getvalue()


def sample_docx():
    b=io.BytesIO()
    doc=Document()
    doc.add_paragraph('Реестровый BIOS требуется в составе оборудования.')
    doc.add_paragraph('Срок поставки: 30 календарных дней с даты заключения контракта.')
    doc.save(b)
    return b.getvalue()


class PackageImport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory(prefix='qx_tender_pkg_')
        cls.original=(store.APP_DATA,store.DB_FILE,store.DOCUMENTS,store._INITIALIZED)
        store.APP_DATA=Path(cls.tmp.name)
        store.DB_FILE=store.APP_DATA/'tenders.sqlite3'
        store.DOCUMENTS=store.APP_DATA/'tender_documents'
        store._INITIALIZED=False
        store.initialize()
        cls.client=TestClient(app)

    @classmethod
    def tearDownClass(cls):
        store.APP_DATA,store.DB_FILE,store.DOCUMENTS,store._INITIALIZED=cls.original
        cls.tmp.cleanup()

    def new(self): return store.create({'title':'Тест ZIP','status':'Черновик'})['id']

    def test_import_document_package_and_bios(self):
        tid=self.new()
        files=[('docs/Проект контракта.docx',sample_docx()),
               ('docs/Извещение.xml',b'<notice><number>0322100002526000466</number></notice>'),
               ('README.txt','Адрес поставки: г. Нижний Новгород'.encode()),
               ('setup.exe',b'MZ fake binary'),('nested/archive.zip',b'nested')]
        response=self.client.post(f'/api/tenders/{tid}/documents/import-package',content=archive(files))
        self.assertEqual(response.status_code,200,response.text)
        parsed=response.json()
        self.assertEqual(len(parsed['imported']),3)
        self.assertEqual(len(parsed['ignored']),2)
        self.assertEqual(store.get(tid)['status'],'Черновик')
        self.assertEqual(len(store.detail(tid)['documents']),3)
        self.assertTrue(any('Импорт пакета' in x['label'] for x in store.detail(tid)['events']))
        analyzed=self.client.post(f'/api/tenders/{tid}/analyze').json()['analysis']
        self.assertIn('Найдено требование',analyzed['bios']['status'])
        self.assertEqual(analyzed['bios']['evidence'][0]['file'],'Проект контракта.docx')

    def test_scanned_pdf_retained_for_manual_review(self):
        from pypdf import PdfWriter
        output=io.BytesIO();writer=PdfWriter();writer.add_blank_page(width=595,height=842);writer.write(output)
        tid=self.new()
        response=self.client.post(f'/api/tenders/{tid}/documents/import-package',content=archive([('скан.pdf',output.getvalue())]))
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(len(response.json()['imported']),1)
        self.assertEqual(len(response.json()['warnings']),1)
        self.assertEqual(store.detail(tid)['documents'][0]['filename'],'скан.pdf')

    def test_duplicate_import_is_idempotent(self):
        tid=self.new()
        zipbytes=archive([('a.txt',b'first document')])
        first=self.client.post(f'/api/tenders/{tid}/documents/import-package',content=zipbytes)
        second=self.client.post(f'/api/tenders/{tid}/documents/import-package',content=zipbytes)
        self.assertEqual(first.status_code,200)
        self.assertEqual(second.status_code,200)
        self.assertEqual(len(second.json()['imported']),0)
        self.assertEqual(second.json()['duplicates'],['a.txt'])
        self.assertEqual(len(store.detail(tid)['documents']),1)

    def test_reject_directory_escape_entire_batch(self):
        tid=self.new()
        before=len(store.detail(tid)['documents'])
        malicious=archive([('ok.txt',b'ok'),('../../evil.txt',b'intruder')])
        res=self.client.post(f'/api/tenders/{tid}/documents/import-package',content=malicious)
        self.assertEqual(res.status_code,400)
        self.assertEqual(before,len(store.detail(tid)['documents']))
        self.assertFalse((store.APP_DATA.parent/'evil.txt').exists())

    def test_reject_zip_bomb_before_storage(self):
        tid=self.new()
        malicious=archive([('very_compressible.txt',b'X'*900000)])
        res=self.client.post(f'/api/tenders/{tid}/documents/import-package',content=malicious)
        self.assertEqual(res.status_code,400)
        self.assertEqual(len(store.detail(tid)['documents']),0)

    def test_reject_corrupt_zip_and_empty_documents(self):
        tid=self.new()
        self.assertEqual(self.client.post(f'/api/tenders/{tid}/documents/import-package',content=b'invalid zip').status_code,400)
        res=self.client.post(f'/api/tenders/{tid}/documents/import-package',content=archive([('bad.docx',b'not docx'),('app.exe',b'MZ')]))
        self.assertEqual(res.status_code,400)
        self.assertEqual(len(store.detail(tid)['documents']),0)

    def test_404_when_tender_absent(self):
        res=self.client.post('/api/tenders/unknown/documents/import-package',content=archive([('a.txt',b'ok')]))
        self.assertEqual(res.status_code,404)

    def test_body_size_limit(self):
        tid=self.new()
        response=self.client.post(f'/api/tenders/{tid}/documents/import-package',content=b'1',headers={'content-length':str(55*1024*1024)})
        self.assertEqual(response.status_code,413)


if __name__=='__main__':unittest.main()
