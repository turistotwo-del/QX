"""Acceptance checks for evidence-only DOCX export and source-only portable launcher."""
import os
import unittest
import tempfile
from pathlib import Path
from io import BytesIO
from docx import Document
from fastapi.testclient import TestClient

from modern import tender_store as t
from modern.server import app
from modern.tender_report import build_tender_report

ROOT = Path(__file__).resolve().parents[1]


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='qx24_report_')
        self.before = (t.APP_DATA, t.DB_FILE, t.DOCUMENTS, t._INITIALIZED)
        t.APP_DATA = Path(self.temp.name)
        t.DB_FILE = t.APP_DATA / 'tenders.sqlite3'
        t.DOCUMENTS = t.APP_DATA / 'tender_documents'
        t._INITIALIZED = False
        t.initialize()
        self.client = TestClient(app)

    def tearDown(self):
        t.APP_DATA, t.DB_FILE, t.DOCUMENTS, t._INITIALIZED = self.before
        self.temp.cleanup()

    def test_word_report_contains_real_source(self):
        row = t.create({'number': '0322100002526000466', 'title': 'Тестовая закупка', 'customer':'Заказчик Тест'})
        d = Document()
        d.add_paragraph('Срок поставки: 30 календарных дней с даты подписания контракта.')
        d.add_paragraph('Поставляемый товар должен иметь реестровый BIOS Минцифры.')
        mem = BytesIO(); d.save(mem)
        t.add_document(row['id'], 'контракт.docx', mem.getvalue(), '')
        response = self.client.get(f"/api/tenders/{row['id']}/report.docx")
        self.assertEqual(response.status_code, 200, response.text[:500] if response.status_code!=200 else '')
        self.assertTrue(response.content.startswith(b'PK'))
        document = Document(BytesIO(response.content))
        text = '\n'.join(p.text for p in document.paragraphs) + '\n' + '\n'.join(c.text for table in document.tables for row in table.rows for c in row.cells)
        self.assertIn('0322100002526000466', text)
        self.assertIn('реестровый BIOS', text)
        self.assertIn('контракт.docx', text)
        self.assertIn('30 календарных дней', text)
        self.assertEqual(t.detail(row['id'])['title'], 'Тестовая закупка')

    def test_report_empty_case_truthful(self):
        row = t.create({'title': 'Без документов'})
        result=self.client.get(f"/api/tenders/{row['id']}/report.docx")
        self.assertEqual(result.status_code,200)
        d=Document(BytesIO(result.content))
        self.assertTrue(any('Документы не приложены' in p.text for p in d.paragraphs))
        self.assertTrue(any('Документы недоступны' in p.text for p in d.paragraphs))

    def test_report_missing_case_404(self):
        self.assertEqual(self.client.get('/api/tenders/missing/report.docx').status_code,404)

    def test_report_missing_attachment_warns_not_certifies(self):
        row=t.create({'title':'С повреждённым файлом'})
        d=Document();d.add_paragraph('BIOS')
        mem=BytesIO();d.save(mem)
        doc_id=t.add_document(row['id'],'тест.docx',mem.getvalue(),'')
        # Simulate missing content after database record was created.
        with t.connection() as conn:
            path=conn.execute('SELECT file_path FROM tender_documents WHERE id=?',(doc_id,)).fetchone()['file_path']
        Path(path).unlink()
        r=self.client.get(f"/api/tenders/{row['id']}/report.docx")
        self.assertEqual(r.status_code,200)
        report=Document(BytesIO(r.content))
        text='\n'.join(p.text for p in report.paragraphs)
        self.assertIn('НЕ ПРОВЕРЕН',text)
        self.assertIn('не все документы',text)


class PortableChecks(unittest.TestCase):
    def test_portable_launcher_only_local_bootstrap(self):
        launcher=(ROOT/'portable'/'start_qx.ps1').read_text(encoding='utf-8')
        self.assertIn('python-3.12.10-embeddable-amd64.zip', launcher)
        self.assertIn('156c7eea90d58cd7e91a23f28a0056616b13e9f4cf4901b7b99b837b7848c6da',launcher)
        self.assertIn('Get-FileHash',launcher)
        self.assertIn('Lib\\site-packages',launcher)
        self.assertIn('--selfcheck',launcher)
        self.assertNotIn('winget install',launcher.lower())
        self.assertNotIn('pip install --user',launcher.lower())
        self.assertTrue((ROOT/'ЗАПУСТИТЬ_QX.cmd').is_file())

    def test_ui_has_report_export(self):
        ts=(ROOT/'modern/web/app.ts').read_text(encoding='utf-8')
        self.assertIn('data-action="tender-export-report"',ts)
        self.assertIn("/report.docx",ts)


if __name__=='__main__':unittest.main()
