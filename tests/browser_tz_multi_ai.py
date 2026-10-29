"""Chromium integration: multi-select, per-document preview, project persistence, AI review."""
import base64
import json
import os
import tempfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

os.environ['QX_OFFLINE_MODE']='1'
os.environ['QX_DATA_DIR']=tempfile.mkdtemp(prefix='qx_tz_multi_')
from docx import Document
from fastapi.testclient import TestClient
from playwright.sync_api import sync_playwright
from chromium_launch import launch_chromium
from modern.server import app
from modern import local_ai, tz_ai

root=Path(__file__).resolve().parents[1]
assets=root/'modern/web'
client=TestClient(app)

def call_api(req):
    path=str(req.get('url','')).replace('/api/','/',1)
    content=base64.b64decode(req['binary']) if req.get('binary') else str(req.get('body') or '').encode('utf-8')
    response=client.request(req.get('method') or 'GET','/api'+path,content=content,headers=req.get('headers') or {})
    return {'status':response.status_code,'body':response.text,'type':response.headers.get('content-type','application/json')}

def sample(s):
    doc=Document();doc.add_heading('Техническое задание',0);doc.add_paragraph(s)
    mem=BytesIO();doc.save(mem)
    return mem.getvalue()

html=(assets/'index.html').read_text('utf-8').replace('<link rel="icon" href="/assets/qx_icon.png">','')
html=html.replace('<link rel="stylesheet" href="/style.css">','<style>'+ (assets/'style.css').read_text('utf-8')+'</style>')
html=html.replace('<script type="module" src="/app.js"></script>', '''<script>
window.fetch=async(url,options={})=>{
let binary=null;
if(options.body instanceof Blob){const a=new Uint8Array(await options.body.arrayBuffer());let x='';for(let i=0;i<a.length;i+=14000)x+=String.fromCharCode(...a.slice(i,i+14000));binary=btoa(x);}
const r=await QXAPI({url:String(url),method:options.method||'GET',body:binary?'':options.body||'',binary,headers:options.headers||{}});
return new Response(r.body,{status:r.status,headers:{'content-type':r.type}});
};
</script><script>'''+(assets/'app.js').read_text('utf-8')+'</script>')


def execute():
    with patch.object(local_ai,'available_models',return_value={'models':['mock-local:1'],'available':True,'message':'Модель запущена.'}), \
         patch.object(tz_ai,'analyze_tz',side_effect=lambda src,req,model: {'model':model,'mode':'Тест модели','warning':'Проверить вручную','rejected':0,'proposals':[{'field':'ram_gb_min','value':64,'previous':req.get('ram_gb_min'),'quote':'Оперативная память не менее 64 ГБ'}]}):
      with sync_playwright() as p:
        browser=launch_chromium(p)
        page=browser.new_page(viewport={'width':1380,'height':850});page.set_default_timeout(9000);print('browser ready',flush=True)
        errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
        page.expose_function('QXAPI',call_api)
        page.set_content(html,wait_until='domcontentloaded')
        print('STEP 0',flush=True);page.wait_for_selector('.sidebar')
        page.locator('#workspace-nav-group > summary').click()
        page.locator('[data-action="tab"][data-tab="tz"]').click()
        page.locator('#file-tz').set_input_files([
            {'name':'first.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
             'buffer':sample('Системный блок. Оперативная память не менее 64 ГБ.')},
            {'name':'second.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
             'buffer':sample('Моноблок. Оперативная память не менее 16 ГБ.')}
        ])
        print('STEP 1',flush=True);page.wait_for_function("document.querySelectorAll('[data-action=tz-doc-view]').length === 2",timeout=15000)
        page.wait_for_function("document.querySelectorAll('.position-chip').length >= 3")
        assert page.evaluate('state.tzDocuments.length') == 2
        assert page.evaluate('state.reqs.every(r=>r._qx_doc_id)')
        assert page.evaluate('state.source_filename') == 'first.docx'
        page.locator('[data-action="tz-doc-view"]').nth(1).click()
        assert page.evaluate('state.source_filename') == 'second.docx'
        page.locator('[data-action="view-source"]').first.click()
        print('STEP 2',flush=True);page.wait_for_selector('.modal')
        page.locator('[data-action="close-modal"]').first.click()
        # Select & preserve a manual price when another file is appended.
        page.locator('[data-action="select"]').first.click()
        print('STEP 3',flush=True);page.wait_for_selector('.price-table')
        page.evaluate("state.plans[0].rows[0].price=12345;state.plans[0].rows[0]._locked=true;")
        page.locator('[data-action="tab"][data-tab="tz"]').click()
        page.locator('#file-tz').set_input_files({'name':'third.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'buffer':sample('Системный блок. Оперативная память не менее 8 ГБ.')})
        print('STEP 4',flush=True);page.wait_for_function("state.tzDocuments.length===3",timeout=18000)
        assert page.evaluate('state.plans[0].rows[0].price') == 12345
        # Run actual UI proposal-review + model-select flow, with mocked inference.
        page.locator('[data-action="tz-ai-models"]').click()
        page.wait_for_selector('#tz-ai-model option[value="mock-local:1"]',state="attached")
        page.locator('#tz-ai-model').select_option('mock-local:1')
        page.locator('[data-action="tz-ai-run"]').click()
        print('STEP 5',flush=True);page.wait_for_selector('[data-tz-ai-check="0"]',timeout=15000)
        assert page.evaluate('state.reqs[state.pos].ram_gb_min') != 64
        page.locator('[data-action="tz-ai-apply"]').click()
        print('STEP 6',flush=True);page.wait_for_selector('.price-table',timeout=15000)
        assert page.evaluate('state.reqs[state.pos].ram_gb_min') == 64
        page.evaluate('window.PROJ=dataPackage()')
        data=page.evaluate('window.PROJ')
        saved=client.post('/api/project/save',json=data)
        assert saved.status_code==200,saved.text
        loaded=client.post('/api/project/load',content=saved.content)
        assert loaded.status_code==200,loaded.text
        page.evaluate('(v)=>importProject(JSON.parse(v))',loaded.text)
        assert page.evaluate('state.tzDocuments.length') == 3
        assert page.evaluate('state.reqs.length') >= 3
        page.locator('[data-action="tab"][data-tab="tz"]').click()
        page.on('dialog',lambda d:d.accept())
        page.locator('[data-action="tz-doc-remove"]').nth(1).click()
        assert page.evaluate('state.tzDocuments.length') == 2
        assert not errors,errors
        print('STEP 7',flush=True);browser.close()
        print('Browser multi TZ + AI review: OK')

if __name__=='__main__':execute()
