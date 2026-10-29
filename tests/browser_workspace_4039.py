"""Workspace acceptance: real API/storage, fixture EIS network, Chromium desktop/mobile.

Run from repository root: PYTHONPATH=. python tests/browser_workspace_4039.py
No live procurement data or external website is used by this deterministic suite.
"""
from __future__ import annotations

import base64
import os
from pathlib import Path
import tempfile
from io import BytesIO
from unittest.mock import patch

os.environ['QX_OFFLINE_MODE'] = '1'
_DATA = tempfile.TemporaryDirectory(prefix='qx_workspace_browser_')
os.environ['QX_DATA_DIR'] = _DATA.name

from fastapi.testclient import TestClient
from docx import Document
from playwright.sync_api import sync_playwright
from chromium_launch import launch_chromium
from modern.server import app
from modern.eis_public import parse_notice_metadata

ROOT = Path(__file__).resolve().parents[1]
NUMBER = '0322100002526000466'
FILE_URL = 'https://zakupki.gov.ru/documents/workspace-specification.docx'
CLIENT = TestClient(app)
REQUESTS: list[dict] = []
VALIDATION_POSTS = 0


def word_document() -> bytes:
    doc = Document()
    doc.add_paragraph('Техническое задание. Поставка компьютеров: количество 10 штук.')
    doc.add_paragraph('Срок поставки 30 календарных дней с даты заключения контракта.')
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def call_api(request):
    path = str(request.get('url', ''))
    REQUESTS.append({'path': path, 'method': request.get('method'), 'body': request.get('body')})
    response = CLIENT.request(request.get('method') or 'GET', path,
                              content=base64.b64decode(request['binary']) if request.get('binary') else str(request.get('body') or '').encode(),
                              headers=request.get('headers') or {})
    mime = response.headers.get('content-type', 'application/json')
    is_binary = 'json' not in mime and not mime.startswith('text/')
    return {'status': response.status_code,
            'body': base64.b64encode(response.content).decode() if is_binary else response.text,
            'binary': is_binary, 'type': mime}


def embedded_app() -> str:
    assets = ROOT / 'modern/web'
    html = (assets / 'index.html').read_text(encoding='utf-8')
    html = html.replace('<link rel="stylesheet" href="/style.css">',
                        '<style>' + (assets / 'style.css').read_text(encoding='utf-8') + '</style>')
    html = html.replace('<script type="module" src="/app.js"></script>',
                        '<script>window.fetch=async(url,options={})=>{'
                        'let binary=null;if(options.body instanceof Blob){'
                        'let bytes=new Uint8Array(await options.body.arrayBuffer()),x="";'
                        'for(let i=0;i<bytes.length;i+=14000)x+=String.fromCharCode(...bytes.slice(i,i+14000));binary=btoa(x);}'
                        'const r=await QXTESTAPI({url:String(url),method:options.method||"GET",'
                        'body:binary?"":options.body||"",binary,headers:options.headers||{}});'
                        'const body=r.binary?Uint8Array.from(atob(r.body),c=>c.charCodeAt(0)):r.body;'
                        'return new Response(body,{status:r.status,headers:{"content-type":r.type}});};</script>'
                        '<script>' + (assets / 'app.js').read_text(encoding='utf-8') + '</script>')
    return html


def wait_import_finished(page, expected_requests):
    page.wait_for_function('!document.querySelector("[data-action=home-eis-import]")?.disabled', timeout=30000)
    page.wait_for_function('typeof state !== "undefined" && state.eisImportBusy === false', timeout=30000)
    assert len([r for r in REQUESTS if r['path'] == '/api/eis/jobs' and r['method'] == 'POST']) - VALIDATION_POSTS == expected_requests


def main():
    global VALIDATION_POSTS
    from modern.eis_discovery import discover_and_import
    xml = (f'<epNotice><purchaseNumber>{NUMBER}</purchaseNumber>'
           '<purchaseName>Поставка рабочих станций</purchaseName>'
           '<customer><fullName>Тестовый заказчик ЕИС</fullName></customer>'
           '<maxPrice>1250000.50</maxPrice><endDate>2026-12-12T12:30:00</endDate></epNotice>').encode()
    metadata = {'available': True, 'notice': parse_notice_metadata(xml, NUMBER),
                'source_url': f'https://zakupki.gov.ru/epz/order/notice/printForm/viewXml.html?regNumber={NUMBER}',
                'search_url': f'https://zakupki.gov.ru/epz/order/extendedsearch/results.html?searchString={NUMBER}',
                'message': 'Поля получены из XML ЕИС.'}
    discovered = {'notice_number': NUMBER, 'documents': [{'url': FILE_URL, 'title': 'ТЗ'}],
                  'found': 1, 'checked_pages': 1, 'available': True, 'complete': True, 'errors': [],
                  'search_url': metadata['search_url'], 'message': 'Найден один документ.'}
    downloads = []
    document_bytes = word_document()  # Reuse identical ZIP bytes; DOCX timestamps vary per creation.

    def download(url, *, opener=None):
        downloads.append(url)
        assert url == FILE_URL
        return 'Техническое задание.docx', document_bytes

    def import_documents(tender_id, **kwargs):
        return discover_and_import(tender_id, discovered=discovered)

    preview = ROOT / 'preview'
    preview.mkdir(exist_ok=True)
    errors = []
    with sync_playwright() as playwright:
        browser = launch_chromium(playwright)
        page = browser.new_page(viewport={'width': 1440, 'height': 1080})
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.expose_function('QXTESTAPI', call_api)
        page.set_content(embedded_app(), wait_until='domcontentloaded')
        number_input = page.locator('#home-eis-number')
        number_input.wait_for()

        # Invalid numbers are rejected locally without a draft or source request.
        number_input.fill('12345')
        button = page.locator('.qx-purchase-form [data-action="home-eis-import"]')
        with patch('modern.eis_sync.fetch_notice_metadata') as rejected_source:
            button.click()
            page.wait_for_function('state.eisImportBusy === false && Boolean(state.eisImportError)')
            assert CLIENT.get('/api/tenders').json()['items'] == []
            rejected_source.assert_not_called()
        VALIDATION_POSTS = len([r for r in REQUESTS if r['path'] == '/api/eis/jobs' and r['method'] == 'POST'])

        with patch('modern.eis_sync.fetch_notice_metadata', return_value=metadata), \
             patch('modern.eis_sync.discover_and_import', side_effect=import_documents), \
             patch('modern.eis_document_download.fetch_official_document', side_effect=download):
            number_input.fill(NUMBER)
            number_input.press('Enter')
            page.wait_for_function('typeof state !== "undefined" && state.eisImportResult?.tender?.id', timeout=30000)
            wait_import_finished(page, 1)
            card_id = page.evaluate('state.eisImportResult.tender.id')
            card = CLIENT.get('/api/tenders/' + card_id).json()['tender']
            assert card['title'] == 'Поставка рабочих станций', card
            assert card['customer'] == 'Тестовый заказчик ЕИС', card
            assert float(card['nmck']) == 1250000.50, card
            assert card['application_due_on'] == '2026-12-12', card
            assert len(card['documents']) == 1, card
            assert card['status'] == 'Черновик'
            assert downloads == [FILE_URL]
            page.screenshot(path=str(preview / 'Workspace_4_0_39_Desktop.png'), full_page=True, animations='disabled')

            # A rerun reuses the exact card and preserves the user's own correction.
            assert CLIENT.patch('/api/tenders/' + card_id, json={'customer': 'Мой заказчик'}).status_code == 200
            # Two same-tick clicks while the request is busy must yield one import.
            page.evaluate('''() => {
                document.querySelector('[data-action="home-eis-import"]').dispatchEvent(new MouseEvent('click', {bubbles:true}));
                document.querySelector('[data-action="home-eis-import"]').dispatchEvent(new MouseEvent('click', {bubbles:true}));
            }''')
            page.wait_for_function('typeof state !== "undefined" && !state.eisImportBusy', timeout=30000)
            wait_import_finished(page, 2)
            cards = CLIENT.get('/api/tenders').json()['items']
            assert len(cards) == 1, cards
            card = CLIENT.get('/api/tenders/' + card_id).json()['tender']
            assert card['customer'] == 'Мой заказчик', card
            assert len(card['documents']) == 1, card['documents']

        # Clicking a result must open its actual tender rather than remain home.
        page.locator(f'[data-action="tender-open"][data-id="{card_id}"]').first.click()
        page.wait_for_selector('[data-tender-panel="overview"]')
        assert page.locator('[data-tender-field="number"]').input_value() == NUMBER
        # Editing two different fields within the debounce window must save both.
        page.evaluate('''() => {
            for (const [field,value] of [['lot','2'],['title','Поставка рабочих станций — проверено']]) {
                const input=document.querySelector('[data-tender-field="'+field+'"]');
                input.value=value;input.dispatchEvent(new Event('change',{bubbles:true}));
            }
        }''')
        page.wait_for_function('!tenderFieldSaves.has("' + card_id + '")', timeout=10000)
        edited = CLIENT.get('/api/tenders/' + card_id).json()['tender']
        assert edited['lot'] == '2', edited
        assert edited['title'] == 'Поставка рабочих станций — проверено', edited
        page.locator('[data-action="tender-detail-tab"][data-view="docs"]').click()
        with patch('modern.eis_public.fetch_notice_metadata', return_value=metadata):
            page.locator('[data-action="tender-use-tz"]').click()
            page.wait_for_function('state.tab === "tz" && !state.busy && state.reqs.length > 0 && state.source_filename === "Техническое задание.docx"', timeout=30000)
            page.wait_for_function('state.tenderEisPreview?.available === true', timeout=30000)
        assert page.evaluate('state.tenderLinkEnabled') is True
        linked_cards = CLIENT.get('/api/tenders').json()['items']
        assert len(linked_cards) == 1, linked_cards
        assert page.evaluate('state.tenderId') == card_id, linked_cards
        assert page.evaluate('state.source_text.includes("Техническое задание")') is True
        page.locator('[data-action="tab"][data-tab="home"]').click()
        page.locator('#home-eis-number').wait_for()

        # An unavailable source produces an honest partial result without erasing data.
        unavailable = {**metadata, 'available': False, 'notice': None,
                       'message': 'ЕИС временно недоступна: проверяемые данные не получены.'}
        page.locator('#home-eis-docs').uncheck()
        with patch('modern.eis_sync.fetch_notice_metadata', return_value=unavailable), \
             patch('modern.eis_sync.fetch_notice_html_metadata', return_value=unavailable), \
             patch('modern.eis_sync.lookup_secondary', return_value={**unavailable, 'documents': [], 'found': 0, 'source_name': 'Вторичный тестовый источник'}):
            page.locator('.qx-purchase-form [data-action="home-eis-import"]').click()
            page.wait_for_function('typeof state !== "undefined" && !state.eisImportBusy', timeout=30000)
            wait_import_finished(page, 3)
        card = CLIENT.get('/api/tenders/' + card_id).json()['tender']
        assert card['customer'] == 'Мой заказчик'
        assert len(card['documents']) == 1
        assert page.evaluate('state.eisImportResult.partial') is True
        assert page.evaluate('state.eisImportResult.metadata_available') is False
        assert 'не предоставили проверяемые сведения' in page.locator('.qx-import-feedback.partial').inner_text()

        # The calculation tools form a native keyboard-operable disclosure group.
        group = page.locator('details').filter(has=page.locator('[data-action="tab"][data-tab="build"]')).first
        assert group.count() == 1
        if group.get_attribute('open') is None:
            group.locator('summary').click()
        page.locator('.sidebar [data-action="tab"][data-tab="build"]').click()
        page.wait_for_function('state.tab === "build"')
        page.locator('[data-action="tab"][data-tab="home"]').click()
        page.locator('#home-eis-number').wait_for()
        page.set_viewport_size({'width': 390, 'height': 844})
        page.wait_for_function('document.querySelectorAll("#toasts .toast").length === 0', timeout=7000)
        page.screenshot(path=str(preview / 'Workspace_4_0_39_Mobile.png'), full_page=True, animations='disabled')
        dimensions = page.evaluate('({width:document.documentElement.clientWidth, scroll:document.documentElement.scrollWidth})')
        assert dimensions['scroll'] <= dimensions['width'] + 1, dimensions
        main_width = page.locator('.main').bounding_box()['width']
        assert main_width >= dimensions['width'] - 80, (main_width, dimensions)
        assert page.locator('#home-eis-number').bounding_box()['width'] >= 160
        sidebar_width = page.locator('.sidebar').bounding_box()['width']
        assert sidebar_width <= 80, page.locator('.sidebar').evaluate('(el) => ({box:el.getBoundingClientRect().width,width:getComputedStyle(el).width,minWidth:getComputedStyle(el).minWidth,maxWidth:getComputedStyle(el).maxWidth})')
        assert not errors, errors
        browser.close()
    print('Workspace 4.0.39 Chromium: success import, preserved fields, deduplication, invalid input, unavailable source, navigation, mobile layout PASS')


if __name__ == '__main__':
    try:
        main()
    finally:
        CLIENT.close()
        _DATA.cleanup()
