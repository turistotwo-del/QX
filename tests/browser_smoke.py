"""Optional Chromium E2E smoke-test without external HTTP access.
Run from repository root: python tests/browser_smoke.py
Requires playwright and a local Chromium executable.
"""
import base64
import os
os.environ['QX_OFFLINE_MODE'] = '1'
import tempfile
from pathlib import Path

os.environ['QX_DATA_DIR']=tempfile.mkdtemp(prefix='qx_browser_e2e_')
from fastapi.testclient import TestClient
from playwright.sync_api import sync_playwright
from chromium_launch import launch_chromium
from modern.server import app
from docx import Document
from io import BytesIO

root=Path(__file__).resolve().parents[1]
assets=root/'modern'/'web'
client=TestClient(app)

def api_call(req):
    path=str(req.get('url','')).replace('/api/','/',1)
    method=req.get('method') or 'GET'
    data=base64.b64decode(req['binary']) if req.get('binary') else str(req.get('body') or '').encode('utf-8')
    rsp=client.request(method,'/api'+path,content=data,headers=req.get('headers') or {})
    return {'status':rsp.status_code,'body':rsp.text,'type':rsp.headers.get('content-type','application/json')}


def test_browser():
    html=(assets/'index.html').read_text(encoding='utf-8')
    html=html.replace('<link rel="icon" href="/assets/qx_icon.png">','')
    html=html.replace('<link rel="stylesheet" href="/style.css">','<style>'+(assets/'style.css').read_text(encoding='utf-8')+'</style>')
    stub='''<script>
    window.fetch=async(url,options={})=>{
      let binary=null;
      if(options.body instanceof Blob){const a=new Uint8Array(await options.body.arrayBuffer());let x='';for(let i=0;i<a.length;i+=15000)x+=String.fromCharCode(...a.slice(i,i+15000));binary=btoa(x);}
      const r=await QXTESTAPI({url:String(url),method:options.method||'GET',body:binary?'':options.body||'',binary,headers:options.headers||{}});
      return new Response(r.body,{status:r.status,headers:{'content-type':r.type}});
    };
    </script>'''
    html=html.replace('<script type="module" src="/app.js"></script>',stub+'<script>'+(assets/'app.js').read_text(encoding='utf-8')+'</script>')
    with sync_playwright() as p:
        import shutil
        chromium=shutil.which('chromium') or shutil.which('chrome')
        browser=launch_chromium(p)
        page=browser.new_page(viewport={'width':1400,'height':900})
        errors=[]
        page.on('pageerror',lambda e: errors.append(str(e)))
        page.expose_function('QXTESTAPI',api_call)
        page.set_content(html,wait_until='domcontentloaded')
        page.wait_for_selector('.sidebar')
        # Real Word document, no pre-existing example DB.
        doc=Document();doc.add_heading('Техническое задание',0)
        doc.add_paragraph('Моноблок с диагональю 27 дюймов. Объём оперативной памяти не менее 16 ГБ.')
        doc.add_paragraph('Реестровый BIOS обязателен.')
        doc.add_paragraph('Срок поставки: 30 календарных дней с даты заключения контракта.')
        out=BytesIO();doc.save(out)
        # QX 4.0.34 requires explicit opt-in, never auto-create a tender.
        # Calculation tools now live in an explicit disclosure group.
        page.locator('#workspace-nav-group > summary').click()
        page.locator('[data-action="tab"][data-tab="tz"]').click()
        page.locator('[data-tender-auto-link]').check()
        page.locator('[data-tender-notice-input]').fill('32500001234')
        page.locator('[data-tender-notice-input]').press('Tab')
        page.locator('#file-tz').set_input_files({'name':'QX_TEST_TZ.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':out.getvalue()})
        page.wait_for_selector('.loaded-source',timeout=12000)
        page.locator('[data-action="select"]').first.click()
        page.wait_for_selector('.price-table',timeout=12000)
        if page.locator('#tender-nav-toggle').get_attribute('aria-expanded') != 'true':
            page.locator('#tender-nav-toggle').click()
        page.locator('[data-action="tender-nav"][data-view="registry"]').click()
        page.wait_for_selector('[data-tender-row]',timeout=12000)
        page.locator('[data-tender-search]').fill('no-such-purchase')
        assert page.locator('[data-tender-row]:visible').count()==0
        assert page.locator('#tender-no-matches:visible').count()==1
        page.locator('[data-tender-search]').fill('')
        assert page.locator('[data-tender-row]:visible').count()==1
        if os.environ.get('QX_SCREENSHOT_DIR'):
            page.wait_for_timeout(450)
            loc=Path(os.environ['QX_SCREENSHOT_DIR']);loc.mkdir(parents=True,exist_ok=True)
            page.screenshot(path=str(loc/'Tender_Registry_Dark.png'),full_page=False)
        page.locator('[data-action="tender-open"]').first.click()
        page.wait_for_selector('.tender-detail-tabs')
        if os.environ.get('QX_SCREENSHOT_DIR'):
            page.wait_for_timeout(450)
            page.screenshot(path=str(Path(os.environ['QX_SCREENSHOT_DIR'])/'Tender_Card_Dark.png'))
        page.locator('[data-action="tender-detail-tab"][data-view="docs"]').click()
        page.wait_for_selector('.tender-doc-list .activity-item',timeout=12000)
        page.locator('[data-action="tender-analyze"]').click()
        page.wait_for_selector('[data-analysis-field="due_days"]',timeout=12000)
        page.locator('[data-analysis-field="due_days"]').check()
        page.locator('[data-action="tender-apply-analysis"]').click()
        page.wait_for_selector('[data-tender-field="due_days"]',timeout=12000)
        assert page.locator('[data-tender-field="due_days"]').input_value()=='30'

        page.locator('[data-action="tender-detail-tab"][data-view="comments"]').click()
        page.locator('#tender-comment').fill('Тестовая заметка')
        page.locator('[data-action="tender-comment"]').click()
        page.wait_for_selector('.tender-comment',timeout=12000)
        page.locator('[data-action="tender-back"]').click()
        page.wait_for_selector('[data-tender-row]',timeout=12000)
        page.locator('[data-action="tender-nav"][data-view="calendar"]').click()
        assert page.locator('.tender-calendar-day').count()>=28
        if os.environ.get('QX_SCREENSHOT_DIR'):
            page.wait_for_timeout(450)
            page.screenshot(path=str(Path(os.environ['QX_SCREENSHOT_DIR'])/'Tender_Calendar_Dark.png'))
        page.locator('[data-action="tender-nav"][data-view="contracts"]').click()
        page.wait_for_selector('h1:has-text("Контракты и поставки")')
        page.locator('[data-action="tender-nav"][data-view="analytics"]').click()
        page.wait_for_selector('h1:has-text("Аналитика тендеров")')
        if os.environ.get('QX_SCREENSHOT_DIR'):
            page.wait_for_timeout(450)
            page.screenshot(path=str(Path(os.environ['QX_SCREENSHOT_DIR'])/'Tender_Analytics_Dark.png'))
        page.locator('[data-action="tender-nav"][data-view="reminders"]').click()
        page.wait_for_selector('h1:has-text("Сроки и напоминания")')
        page.wait_for_selector('[data-tender-notifications]')
        page.locator('[data-tender-notifications]').check()
        page.wait_for_timeout(300)
        assert page.locator('[data-tender-notifications]').is_checked()
        if os.environ.get('QX_SCREENSHOT_DIR'):
            page.screenshot(path=str(Path(os.environ['QX_SCREENSHOT_DIR'])/'Tender_Reminders_Dark.png'))
        page.locator('[data-action="tender-nav"][data-view="organizations"]').click()
        page.wait_for_selector('h1:has-text("Организации")')
        page.locator('#new-org-name').fill('ООО Тест QX')
        page.locator('#new-org-inn').fill('1234567890')
        page.locator('[data-action="tender-add-org"]').click()
        page.wait_for_selector('[data-org-editor]')
        page.locator('[data-org-editor] [data-org-name]').fill('ООО Тест QX обновлено')
        page.locator('[data-action="tender-org-save"]').click()
        page.wait_for_selector('[data-org-name][value="ООО Тест QX обновлено"]')
        if page.locator('#tender-nav-toggle').get_attribute('aria-expanded') != 'true':
            page.locator('#tender-nav-toggle').click()
        page.locator('[data-action="tender-nav"][data-view="registry"]').click()
        page.wait_for_selector('[data-tender-row]')
        assert page.locator('.sidebar .tender-group').count()==1
        assert page.locator('.sidebar .tender-subnav').count()==5
        assert page.locator('.sidebar .qx-calendar-nav').count()==1
        # Both themes must retain tender navigation and readable tables.
        page.evaluate("state.theme='dark';render()")
        assert not page.locator('html').evaluate('(el)=>el.classList.contains("light")')
        page.wait_for_selector('[data-tender-row]')
        if os.environ.get('QX_SCREENSHOT_DIR'):
            page.wait_for_timeout(450)
            page.screenshot(path=str(Path(os.environ['QX_SCREENSHOT_DIR'])/'Tender_Registry_Dark.png'))
        page.evaluate("state.theme='light';render()")
        assert page.locator('html').evaluate('(el)=>el.classList.contains("light")')
        page.wait_for_selector('[data-tender-row]')
        if os.environ.get('QX_SCREENSHOT_DIR'):
            page.wait_for_timeout(450)
            page.screenshot(path=str(Path(os.environ['QX_SCREENSHOT_DIR'])/'Tender_Registry_Light.png'))
        page.locator('[data-action="tab"][data-tab="settings"]').click()
        page.wait_for_selector('[data-setting="offline_mode"]')
        page.locator('[data-setting="offline_mode"]').check()
        page.locator('[data-action="save-settings"]').click()
        page.wait_for_function('!document.querySelector("#busy-overlay")')
        assert client.get('/api/boot').json()['settings']['offline_mode'] is True
        assert not errors, errors
        browser.close()
    print('Browser E2E: OK')

if __name__=='__main__':
    try: test_browser()
    finally: client.close()
